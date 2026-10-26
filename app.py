from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, session, g, abort, Response
from werkzeug.exceptions import HTTPException
from werkzeug.security import generate_password_hash, check_password_hash
import sqlite3
from datetime import date, datetime, timedelta
from functools import wraps
import secrets
import os
import re
import json
import hashlib
import urllib.parse
import urllib.request
import urllib.error
import threading
import time

import backup
import clock
import deep_clean
import gzip
import integration_api
import safety
import settings
import totp
import night
import notifications
import timesheet
import version as app_version_info
import room_state

APP_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("MOTEL_DB_PATH", os.path.join(APP_DIR, "motel.db"))
AUDIT_FILE = os.environ.get("MOTEL_AUDIT_FILE", os.path.join(APP_DIR, "logs", "audit.log"))
USER_COLORS = ["#176b56", "#255f9e", "#a65d16", "#8b4fa3", "#b33f55", "#287f8b", "#6a7d1c", "#5b6470", "#b06c2c", "#3563b5", "#8a3d67", "#39734d"]
def load_secret_key():
    """Cheia de semnare a sesiunilor: din MOTEL_SECRET_KEY sau, dacă lipsește, una aleatorie persistată (nu există cheie implicită)."""
    env = os.environ.get("MOTEL_SECRET_KEY", "")
    if len(env) >= 32:
        return env
    if env:
        print("AVERTISMENT: MOTEL_SECRET_KEY are sub 32 de caractere și este ignorată; se folosește o cheie aleatorie persistată.", file=__import__("sys").stderr)
    path = os.path.join(os.path.dirname(DB_PATH) or ".", "secret.key")
    try:
        with open(path, "r", encoding="ascii") as handle:
            key = handle.read().strip()
        if len(key) >= 32:
            return key
    except OSError:
        pass
    key = secrets.token_hex(32)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, "w", encoding="ascii") as handle:
        handle.write(key)
    return key


app = Flask(__name__)
app.secret_key = load_secret_key()
SECURE_COOKIES = os.environ.get("MOTEL_COOKIE_SECURE", "0") == "1" or bool(os.environ.get("MOTEL_SSL_CERT"))
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax", SESSION_COOKIE_SECURE=SECURE_COOKIES,
                  PERMANENT_SESSION_LIFETIME=timedelta(hours=8), MAX_CONTENT_LENGTH=256 * 1024)
DUMMY_PIN_HASH = generate_password_hash(secrets.token_hex(8))   # egalizează timpul de răspuns pentru utilizatori inexistenți

STATUSES = {"provisional", "confirmed", "unconfirmed", "checked_in", "checked_out", "cancelled"}

def connect():
    conn = sqlite3.connect(DB_PATH, timeout=10, isolation_level=None)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    return conn

def backup_database_before_migration():
    """Create a SQLite-consistent snapshot once before the first schema upgrade."""
    if not os.path.exists(DB_PATH) or os.path.getsize(DB_PATH) == 0:
        return
    try:
        source = sqlite3.connect(DB_PATH, timeout=10)
        version = source.execute("PRAGMA user_version").fetchone()[0]
        has_tables = source.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name IN ('rooms','bookings') LIMIT 1").fetchone()
        if version < 13 and has_tables:
            backup_dir = os.path.join(os.path.dirname(DB_PATH), "backups")
            os.makedirs(backup_dir, exist_ok=True)
            stamp = clock.now().strftime("%Y%m%d_%H%M%S")
            target_path = os.path.join(backup_dir, f"motel.db.before_v{version}_{stamp}.sqlite3")
            target = sqlite3.connect(target_path)
            source.backup(target)
            target.close()
        source.close()
    except sqlite3.Error:
        # Fail closed on a backup error: startup must not proceed with a schema change.
        raise RuntimeError("Nu s-a putut crea copia de siguranță înainte de actualizarea bazei de date.")

def init_db():
    backup_database_before_migration()
    conn = connect()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS rooms (
            id INTEGER PRIMARY KEY,
            name TEXT NOT NULL UNIQUE,
            room_type TEXT NOT NULL DEFAULT 'Dublă',
            building TEXT NOT NULL DEFAULT 'Corp A',
            bathroom_type TEXT NOT NULL DEFAULT 'Baie comună',
            bed_configuration TEXT NOT NULL DEFAULT '2 paturi',
            sort_order INTEGER NOT NULL DEFAULT 0,
            active INTEGER NOT NULL DEFAULT 1,
            out_of_order INTEGER NOT NULL DEFAULT 0,
            maintenance_note TEXT NOT NULL DEFAULT ''
        )
    """)
    # Upgrade a database created by an earlier MVP version, without deleting bookings.
    existing_columns = {row['name'] for row in conn.execute("PRAGMA table_info(rooms)").fetchall()}
    for column, declaration in [
        ('building', "TEXT NOT NULL DEFAULT 'Corp A'"),
        ('bathroom_type', "TEXT NOT NULL DEFAULT 'Baie comună'"),
        ('bed_configuration', "TEXT NOT NULL DEFAULT '2 paturi'"),
        ('sort_order', "INTEGER NOT NULL DEFAULT 0"),
        ('out_of_order', "INTEGER NOT NULL DEFAULT 0"),
        ('maintenance_note', "TEXT NOT NULL DEFAULT ''"),
        ('default_rate', "REAL NOT NULL DEFAULT 0"),   # v12: tarif implicit pe noapte (cazări la recepția de noapte)
    ]:
        if column not in existing_columns:
            conn.execute(f"ALTER TABLE rooms ADD COLUMN {column} {declaration}")
    conn.execute("""
        CREATE TABLE IF NOT EXISTS bookings (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            guest_name TEXT NOT NULL,
            phone TEXT DEFAULT '',
            email TEXT DEFAULT '',
            room_id INTEGER NOT NULL REFERENCES rooms(id),
            check_in TEXT NOT NULL,
            check_out TEXT NOT NULL,
            guests INTEGER NOT NULL DEFAULT 1,
            rate REAL NOT NULL DEFAULT 0,
            deposit REAL NOT NULL DEFAULT 0,
            source TEXT NOT NULL DEFAULT 'Direct',
            status TEXT NOT NULL DEFAULT 'confirmed',
            notes TEXT DEFAULT '',
            created_by INTEGER,
            created_by_name TEXT NOT NULL DEFAULT 'Importat / necunoscut',
            created_by_color TEXT NOT NULL DEFAULT '#71817c',
            group_id TEXT,
            receipt_number TEXT NOT NULL DEFAULT '',
            reconfirmed_at TEXT,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            CHECK (check_out > check_in),
            CHECK (guests > 0),
            CHECK (rate >= 0),
            CHECK (deposit >= 0)
        )
    """)
    booking_columns = {row['name'] for row in conn.execute("PRAGMA table_info(bookings)").fetchall()}
    for column, declaration in [
        ('created_by', "INTEGER"),
        ('created_by_name', "TEXT NOT NULL DEFAULT 'Importat / necunoscut'"),
        ('created_by_color', "TEXT NOT NULL DEFAULT '#71817c'"),
        ('group_id', "TEXT"),
        ('receipt_number', "TEXT NOT NULL DEFAULT ''"),
        ('reconfirmed_at', "TEXT"),
        ('late_arrival', "INTEGER NOT NULL DEFAULT 0"),   # v12: sosește după ora 21:00 (recepție de noapte)
        ('night_note', "TEXT NOT NULL DEFAULT ''"),       # observații de la recepția de zi pentru recepția de noapte
        ('night_payment', "TEXT"),                        # cash | morning | collected (cazări făcute noaptea)
        ('night_by', "INTEGER"),
        ('night_at', "TEXT"),
    ]:
        if column not in booking_columns:
            conn.execute(f"ALTER TABLE bookings ADD COLUMN {column} {declaration}")

    # Seed the exact 25-room inventory supplied by the property. Do not overwrite
    # an existing inventory or any existing reservations.
    count = conn.execute("SELECT COUNT(*) FROM rooms").fetchone()[0]
    if count == 0:
        inventory = []
        order = 0
        corp_a = [
            ('101', '2 paturi'), ('102', '2 paturi'), ('103', '2 paturi'),
            ('104', '1 pat matrimonial'), ('201', '1 pat matrimonial'),
            ('202', '1 pat matrimonial'), ('203', '2 paturi'), ('204', '2 paturi'),
            ('205', '1 pat matrimonial'), ('206', '1 pat matrimonial'),
        ]
        corp_b = [('105', '1 pat matrimonial'), ('106', '1 pat matrimonial'),
                  ('107', '1 pat matrimonial'), ('207', '1 pat matrimonial'),
                  ('208', '1 pat matrimonial'), ('209', '1 pat matrimonial')]
        corp_c = [('01', '2 paturi'), ('02', '2 paturi'), ('03', '2 paturi'),
                  ('108', '2 paturi'), ('109', '2 paturi'), ('110', '2 paturi'),
                  ('211', '2 paturi'), ('212', '2 paturi'), ('213', '2 paturi')]
        for name, beds in corp_a:
            order += 1
            inventory.append((name, 'Cameră cu ' + beds, 'Corp A', 'Baie comună', beds, order))
        for name, beds in corp_b:
            order += 1
            inventory.append((name, 'Cameră cu ' + beds, 'Corp B', 'Baie proprie', beds, order))
        for name, beds in corp_c:
            order += 1
            inventory.append((name, 'Cameră cu ' + beds, 'Corp C', 'Baie proprie', beds, order))
        conn.executemany(
            "INSERT INTO rooms(name, room_type, building, bathroom_type, bed_configuration, sort_order) VALUES (?, ?, ?, ?, ?, ?)",
            inventory
        )
    else:
        # For a previous generic 25-room seed, replace it only if there are no bookings.
        bookings_count = conn.execute("SELECT COUNT(*) FROM bookings").fetchone()[0]
        generic_names = {f"{i:02d}" for i in range(1, 26)}
        current_names = {row['name'] for row in conn.execute("SELECT name FROM rooms").fetchall()}
        if bookings_count == 0 and current_names == generic_names:
            conn.execute("DELETE FROM rooms")
            inventory = []
            order = 0
            for building, bathroom, rows in [
                ('Corp A', 'Baie comună', [('101','2 paturi'),('102','2 paturi'),('103','2 paturi'),('104','1 pat matrimonial'),('201','1 pat matrimonial'),('202','1 pat matrimonial'),('203','2 paturi'),('204','2 paturi'),('205','1 pat matrimonial'),('206','1 pat matrimonial')]),
                ('Corp B', 'Baie proprie', [('105','1 pat matrimonial'),('106','1 pat matrimonial'),('107','1 pat matrimonial'),('207','1 pat matrimonial'),('208','1 pat matrimonial'),('209','1 pat matrimonial')]),
                ('Corp C', 'Baie proprie', [('01','2 paturi'),('02','2 paturi'),('03','2 paturi'),('108','2 paturi'),('109','2 paturi'),('110','2 paturi'),('211','2 paturi'),('212','2 paturi'),('213','2 paturi')]),
            ]:
                for name, beds in rows:
                    order += 1
                    inventory.append((name, 'Cameră cu ' + beds, building, bathroom, beds, order))
            conn.executemany("INSERT INTO rooms(name, room_type, building, bathroom_type, bed_configuration, sort_order) VALUES (?, ?, ?, ?, ?, ?)", inventory)
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT NOT NULL UNIQUE,
            display_name TEXT NOT NULL, pin_hash TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('admin','reception')),
            active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            last_login_at TEXT
        );
        CREATE TABLE IF NOT EXISTS audit_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, username TEXT NOT NULL DEFAULT 'system',
            action TEXT NOT NULL, entity_type TEXT NOT NULL, entity_id TEXT, details TEXT NOT NULL DEFAULT '{}',
            ip_address TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS auth_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT, ip_address TEXT NOT NULL, attempted_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, success INTEGER NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS invoices (
            id INTEGER PRIMARY KEY AUTOINCREMENT, booking_id INTEGER NOT NULL UNIQUE REFERENCES bookings(id),
            provider TEXT NOT NULL DEFAULT 'oblio', mode TEXT NOT NULL DEFAULT 'test', status TEXT NOT NULL,
            idempotency_key TEXT NOT NULL UNIQUE, series_name TEXT, invoice_number TEXT, issue_date TEXT,
            total REAL NOT NULL DEFAULT 0, provider_link TEXT, response_json TEXT NOT NULL DEFAULT '{}',
            created_by INTEGER, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(created_by) REFERENCES users(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS invoice_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT, invoice_id INTEGER REFERENCES invoices(id) ON DELETE SET NULL,
            booking_id INTEGER NOT NULL REFERENCES bookings(id), idempotency_key TEXT NOT NULL, status TEXT NOT NULL,
            message TEXT NOT NULL DEFAULT '', response_json TEXT NOT NULL DEFAULT '{}',
            attempted_by INTEGER, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(attempted_by) REFERENCES users(id) ON DELETE SET NULL
        );
        CREATE INDEX IF NOT EXISTS idx_audit_created_at ON audit_log(created_at);
        CREATE INDEX IF NOT EXISTS idx_invoice_attempts_booking ON invoice_attempts(booking_id, created_at);
    """)
    # Version the schema so the automatic pre-migration backup is not repeated each launch.
    conn.execute("CREATE INDEX IF NOT EXISTS idx_bookings_group_id ON bookings(group_id)")
    conn.execute("""
        CREATE TABLE IF NOT EXISTS room_cleaning (
            room_id INTEGER NOT NULL REFERENCES rooms(id),
            service_date TEXT NOT NULL,
            is_clean INTEGER NOT NULL DEFAULT 0,
            updated_by INTEGER,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY(room_id, service_date)
        )
    """)
    # V12: keep the existing role constraint intact and add a separate access profile
    # so older databases can create a housekeeper account without rebuilding users/FKs.
    user_columns = {row['name'] for row in conn.execute("PRAGMA table_info(users)").fetchall()}
    attempt_columns = {row['name'] for row in conn.execute("PRAGMA table_info(auth_attempts)").fetchall()}
    if 'username' not in attempt_columns:
        conn.execute("ALTER TABLE auth_attempts ADD COLUMN username TEXT NOT NULL DEFAULT ''")
    for column, declaration in [('totp_secret', "TEXT"), ('totp_enabled', "INTEGER NOT NULL DEFAULT 0"), ('totp_last_step', "INTEGER"),
                                ('totp_recovery', "TEXT NOT NULL DEFAULT '[]'"), ('session_epoch', "INTEGER NOT NULL DEFAULT 0")]:
        if column not in user_columns:
            conn.execute(f"ALTER TABLE users ADD COLUMN {column} {declaration}")
    if 'staff_type' not in user_columns:
        conn.execute("ALTER TABLE users ADD COLUMN staff_type TEXT NOT NULL DEFAULT 'reception'")
    cleaning_columns = {row['name'] for row in conn.execute("PRAGMA table_info(room_cleaning)").fetchall()}
    if 'updated_by_name' not in cleaning_columns:
        conn.execute("ALTER TABLE room_cleaning ADD COLUMN updated_by_name TEXT NOT NULL DEFAULT ''")
    conn.execute("""
        CREATE TABLE IF NOT EXISTS housekeeping_requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            room_id INTEGER NOT NULL REFERENCES rooms(id),
            service_date TEXT NOT NULL,
            request_type TEXT NOT NULL,
            note TEXT NOT NULL DEFAULT '',
            status TEXT NOT NULL DEFAULT 'open' CHECK(status IN ('open','done','cancelled')),
            requested_by INTEGER,
            requested_by_name TEXT NOT NULL DEFAULT '',
            completed_by INTEGER,
            completed_by_name TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            completed_at TEXT,
            FOREIGN KEY(requested_by) REFERENCES users(id) ON DELETE SET NULL,
            FOREIGN KEY(completed_by) REFERENCES users(id) ON DELETE SET NULL
        )
    """)
    conn.execute("CREATE INDEX IF NOT EXISTS idx_housekeeping_requests_day_status ON housekeeping_requests(service_date, status)")
    room_state.migrate(conn)  # v9: cod stabil pe cameră + coadă de evenimente pentru integrări
    deep_clean.migrate_request_types(conn)  # v10: hârtie igienică și solicitări combinate
    deep_clean.migrate(conn)                # v10: formular săptămânal de curățenie extinsă
    timesheet.migrate(conn)                 # v11: pontaj cameriste
    settings.migrate(conn)                  # v13: module și ore setate de administrator
    if conn.execute("PRAGMA user_version").fetchone()[0] < 13:
        # v13: tarif de noapte implicit pe corp, doar acolo unde nu era setat (Corp A 150 lei, Corp B/C 180 lei)
        conn.execute("UPDATE rooms SET default_rate = CASE WHEN building='Corp A' THEN 150 ELSE 180 END WHERE default_rate IS NULL OR default_rate <= 0")
    conn.execute("PRAGMA user_version = 13")
    # Keep existing rooms ordered and labeled with their current metadata.
    conn.close()

def overlaps(conn, room_id, check_in, check_out, exclude_id=None):
    sql = """
        SELECT id, guest_name, check_in, check_out
        FROM bookings
        WHERE room_id = ?
          AND status NOT IN ('cancelled', 'unconfirmed')
          AND check_in < ?
          AND check_out > ?
    """
    params = [room_id, check_out, check_in]
    if exclude_id is not None:
        sql += " AND id != ?"
        params.append(exclude_id)
    return conn.execute(sql, params).fetchone()

def apply_direct_booking_cutoff(now=None):
    """At/after 20:30 local server time, release un-reconfirmed same-day direct bookings."""
    now = now or room_state.now_fn().astimezone()   # ceas injectabil (teste); în producție este ora locală
    if not settings.enabled("mod_reconfirm") or now.time().replace(tzinfo=None) < settings.time_value("reconfirm_time"):
        return 0
    today = now.date().isoformat()
    conn = connect()
    rows = conn.execute("""
        SELECT id, guest_name FROM bookings
        WHERE check_in = ? AND source = 'Direct' AND group_id IS NULL
          AND status IN ('confirmed', 'provisional') AND reconfirmed_at IS NULL
          AND date(created_at, 'localtime') = ?
    """, (today, today)).fetchall()
    if rows:
        conn.executemany("UPDATE bookings SET status='unconfirmed', updated_at=CURRENT_TIMESTAMP WHERE id=?", [(r['id'],) for r in rows])
        for row in rows:
            conn.execute("INSERT INTO audit_log(user_id, username, action, entity_type, entity_id, details, ip_address, created_at) VALUES (NULL, 'system', 'automatic_unconfirmed', 'booking', ?, ?, '', ?)",
                         (str(row['id']), json.dumps({"reason":"direct_same_day_not_reconfirmed_by_20_30", "guest_name":row['guest_name']}, ensure_ascii=False), now.isoformat(timespec='seconds')))
        for row in rows:
            try:
                room_state.emit_for_audit(conn, 'automatic_unconfirmed', 'booking', row['id'], {})
            except Exception:
                app.logger.exception("Nu s-a putut înregistra evenimentul de cameră.")
        conn.commit()
    conn.close()
    if rows:
        for row in rows:
            # Background job has no request context; write a compact file log and DB audit entry separately.
            try:
                os.makedirs(os.path.dirname(AUDIT_FILE) or '.', exist_ok=True)
                with open(AUDIT_FILE, 'a', encoding='utf-8') as stream:
                    stream.write(json.dumps({"timestamp": now.isoformat(timespec="seconds"), "username": "system", "action": "automatic_unconfirmed", "entity_type": "booking", "entity_id": row['id'], "details": {"reason": "direct_same_day_not_reconfirmed_by_20_30"}}, ensure_ascii=False) + "\n")
            except OSError:
                app.logger.exception("Nu s-a putut scrie jurnalul pentru schimbarea automată a statusului.")
    return len(rows)


def apply_late_arrival_confirmation():
    """După încheierea programului recepției de zi (21:00), rezervările neconfirmate cu sosire azi devin „vin după ora 21”."""
    now = room_state.now_fn()
    if not settings.enabled("mod_late_confirm") or now.time() < night.reception_close():
        return 0
    conn = connect()
    ids = night.confirm_late_unconfirmed(conn, now.date().isoformat())
    for booking_id in ids:
        conn.execute("INSERT INTO audit_log(user_id, username, action, entity_type, entity_id, details, ip_address, created_at) VALUES (NULL, 'system', 'automatic_late_confirm', 'booking', ?, ?, '', ?)",
                     (str(booking_id), json.dumps({"reason": "unconfirmed_at_reception_close", "status": "confirmed"}), now.isoformat(timespec="seconds")))
        try:
            room_state.emit_for_audit(conn, "automatic_late_confirm", "booking", booking_id, {"status": "confirmed"})
        except Exception:
            app.logger.exception("Nu s-a putut înregistra evenimentul de cameră.")
    conn.commit()
    conn.close()
    return len(ids)


def track_housekeeper_presence():
    """Pontaj: o actualizare la cel mult 30 s; o pauză lungă începe o perioadă nouă."""
    if not settings.enabled("mod_timesheet") or (session.get("ts_id") and time.time() - session.get("ts_touch", 0) < 30):
        return
    try:
        conn = connect()
        session["ts_id"] = timesheet.touch_or_start(conn, session.get("ts_id"), g.user)
        conn.close()
        session["ts_touch"] = time.time()
    except Exception:
        app.logger.exception("Nu s-a putut actualiza pontajul.")


_last_time_events = 0.0


def maybe_emit_time_events():
    """Evenimente generate de trecerea timpului (ex. cameră eliberată automat la 13:00), cel mult o dată la 30 s."""
    global _last_time_events
    if time.monotonic() - _last_time_events < 30:
        return
    _last_time_events = time.monotonic()
    try:
        conn = connect()
        room_state.emit_time_events(conn)
        conn.close()
        apply_late_arrival_confirmation()
    except Exception:
        app.logger.exception("Nu s-au putut genera evenimentele bazate pe oră.")


app.jinja_env.filters["service_label"] = notifications.service_label


@app.template_filter("lei")
def lei_filter(value):
    """1250.5 -> '1.250,50 lei'"""
    try:
        text = f"{float(value):,.2f}"
    except (TypeError, ValueError):
        return "0,00 lei"
    return text.replace(",", "X").replace(".", ",").replace("X", ".") + " lei"


@app.template_filter("dmy")
def dmy_filter(value):
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").strftime("%d.%m")
    except ValueError:
        return value or ""


@app.after_request
def security_headers(response):
    """Antete de securitate: CSP strict cu nonce, anti-clickjacking, nosniff, referrer, permisiuni, HSTS (HTTPS), fără cache pentru pagini."""
    secure = request.is_secure or app.config["SESSION_COOKIE_SECURE"]
    nonce = getattr(g, "csp_nonce", "")
    policy = ["default-src 'self'", f"script-src 'self' 'nonce-{nonce}'", "style-src 'self' 'unsafe-inline'", "img-src 'self' data:",
              "font-src 'self'", "connect-src 'self'", "media-src 'self'", "manifest-src 'self'", "object-src 'none'",
              "base-uri 'self'", "form-action 'self'", "frame-ancestors 'none'"]
    if secure:
        policy.append("upgrade-insecure-requests")
    headers = response.headers
    headers["Content-Security-Policy"] = "; ".join(policy)
    headers["X-Content-Type-Options"] = "nosniff"
    headers["X-Frame-Options"] = "DENY"
    headers["Referrer-Policy"] = "same-origin"
    headers["Cross-Origin-Opener-Policy"] = "same-origin"
    headers["Cross-Origin-Resource-Policy"] = "same-origin"
    headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=(), payment=(), usb=(), interest-cohort=()"
    if secure:
        headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    if request.endpoint != "static" and "Cache-Control" not in headers:
        headers["Cache-Control"] = "no-store"
    headers["Server"] = "PMS"
    return response


ERROR_TEXT = {400: "Cererea nu a putut fi procesată.", 403: "Nu aveți dreptul să accesați această pagină.", 404: "Pagina nu există.",
              405: "Acțiune nepermisă.", 413: "Datele trimise sunt prea mari.", 429: "Prea multe cereri. Reîncercați mai târziu."}


@app.errorhandler(HTTPException)
def http_error(err):
    code = err.code or 500
    if request.path.startswith("/api/"):
        return jsonify({"error": err.name}), code
    # Mesajul propriu (abort(..., description=...)) se afișează; altfel, text generic fără detalii tehnice.
    custom = err.description if isinstance(err.description, str) and err.description != type(err).description else None
    if code == 403:
        app.logger.warning("Acces refuzat: %s %s utilizator=%s", request.method, request.path, g.user["username"] if getattr(g, "user", None) else "-")
    return render_template("error.html", code=code, message=custom or ERROR_TEXT.get(code, "A apărut o problemă.")), code


@app.errorhandler(Exception)
def unexpected_error(err):
    app.logger.exception("Eroare neașteptată la %s %s", request.method, request.path)
    if request.path.startswith("/api/"):
        return jsonify({"error": "Eroare internă"}), 500
    return render_template("error.html", code=500, message="A apărut o eroare internă. Încercați din nou; dacă se repetă, anunțați administratorul."), 500


COMPRESSIBLE = {"text/html", "text/css", "application/javascript", "text/javascript", "application/json", "image/svg+xml", "application/manifest+json"}


@app.after_request
def optimize_response(response):
    """Cache pentru fișiere statice versionate și compresie gzip: pagina de calendar scade de la ~130 KB la ~10 KB pe Wi-Fi."""
    if request.endpoint == "static" and response.status_code == 200:
        response.headers["Cache-Control"] = "public, max-age=604800" if request.args.get("v") else "no-cache"
    if (response.status_code == 200 and response.mimetype in COMPRESSIBLE and "Content-Encoding" not in response.headers
            and "gzip" in request.headers.get("Accept-Encoding", "")):
        response.direct_passthrough = False
        data = response.get_data()
        if len(data) > 1024:
            response.set_data(gzip.compress(data, compresslevel=6))
            response.headers["Content-Encoding"] = "gzip"
            response.headers.add("Vary", "Accept-Encoding")
    return response


def maybe_auto_backup(now=None):
    """Copia zilnică automată (din planificatorul aplicației)."""
    now = now or clock.now()
    dest = backup.folder(DB_PATH)
    if not backup.needed(dest, now, settings.time_value("backup_time")):
        return None
    result = backup.run(DB_PATH, dest, settings.number("backup_keep_days"), now)
    conn = connect()
    conn.execute("INSERT INTO audit_log(user_id, username, action, entity_type, entity_id, details, ip_address, created_at) VALUES (NULL, 'system', 'backup_auto', 'backup', NULL, ?, '', ?)",
                 (json.dumps({"file": os.path.basename(result["path"]), "bookings": result["bookings"], "removed": result["removed"]}), now.isoformat(timespec="seconds")))
    conn.close()
    return result


def cutoff_worker():
    """Local-test scheduler; checks the 20:30 policy every 30 seconds."""
    while True:
        try:
            apply_direct_booking_cutoff()
            maybe_emit_time_events()
            maybe_auto_backup()
        except Exception:
            app.logger.exception("Eroare la aplicarea procedurii automate de la 20:30.")
        time.sleep(30)


def audit(action, entity_type, entity_id=None, details=None):
    """Write audit events to SQLite and an append-only, UTF-8 JSONL file."""
    user = getattr(g, "user", None)
    event = {
        "timestamp": clock.now().astimezone().isoformat(timespec="seconds"),
        "user_id": user["id"] if user else None,
        "username": user["username"] if user else "system",
        "display_name": user["display_name"] if user else "Sistem",
        "action": action,
        "entity_type": entity_type,
        "entity_id": str(entity_id) if entity_id is not None else None,
        "details": details or {},
        "ip_address": request.remote_addr or "",
    }
    conn = connect()
    conn.execute(
        "INSERT INTO audit_log(user_id, username, action, entity_type, entity_id, details, ip_address, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        (event["user_id"], event["username"], action, entity_type, event["entity_id"],
         json.dumps(event["details"], ensure_ascii=False, sort_keys=True), event["ip_address"], event["timestamp"])
    )
    try:
        actor = {"id": user["id"], "name": user["display_name"]} if user else None
        room_state.emit_for_audit(conn, action, entity_type, entity_id, event["details"], actor)
    except Exception:
        # O problemă la coada de evenimente nu trebuie să blocheze operațiunea de la recepție.
        app.logger.exception("Nu s-a putut înregistra evenimentul de cameră.")
    conn.close()
    try:
        os.makedirs(os.path.dirname(AUDIT_FILE) or ".", exist_ok=True)
        with open(AUDIT_FILE, "a", encoding="utf-8") as stream:
            stream.write(json.dumps(event, ensure_ascii=False, sort_keys=True) + "\n")
            stream.flush()
            os.fsync(stream.fileno())
    except OSError:
        # Keep the application operational if the file system is temporarily read-only;
        # the database audit record is still preserved for later export.
        app.logger.exception("Nu s-a putut scrie evenimentul în fișierul de audit.")


def color_for_user(user_id):
    try:
        return USER_COLORS[(int(user_id) - 1) % len(USER_COLORS)]
    except (TypeError, ValueError):
        return "#71817c"

def current_user():
    user_id = session.get("user_id")
    if not user_id:
        return None
    conn = connect()
    user = conn.execute("SELECT id, username, display_name, role, active, COALESCE(staff_type, 'reception') AS staff_type, session_epoch, totp_enabled FROM users WHERE id=?", (user_id,)).fetchone()
    conn.close()
    if not user or not user["active"] or session.get("epoch", 0) != user["session_epoch"]:
        session.clear()          # cont dezactivat, PIN/rol schimbat sau „deconectează toate sesiunile”
        return None
    return user

def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not getattr(g, "user", None):
            return redirect(url_for("login", next=request.full_path if request.query_string else request.path))
        return fn(*args, **kwargs)
    return wrapped

def admin_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not getattr(g, "user", None):
            return redirect(url_for("login", next=request.path))
        if g.user["role"] != "admin":
            abort(403)
        return fn(*args, **kwargs)
    wrapped.admin_only = True      # folosit de testele de securitate pentru a verifica automat toate rutele de administrare
    return wrapped

PUBLIC_ENDPOINTS = {"setup", "login", "static", "login_2fa", "login_2fa_setup"}
PASSIVE_ENDPOINTS = {"notifications_feed", "housekeeping_updates", "static"}   # interogările automate nu prelungesc sesiunea


def idle_limit_seconds(user):
    """Expirare pentru inactivitate: administrator 30 min, restul personalului 12 h (MOTEL_IDLE_ADMIN_MIN / MOTEL_IDLE_STAFF_MIN)."""
    try:
        admin_min = int(os.environ.get("MOTEL_IDLE_ADMIN_MIN", "30"))
        staff_min = int(os.environ.get("MOTEL_IDLE_STAFF_MIN", "720"))
    except ValueError:
        admin_min, staff_min = 30, 720
    return (admin_min if user["role"] == "admin" else staff_min) * 60


@app.before_request
def security_and_authentication():
    g.csp_nonce = secrets.token_urlsafe(16)
    apply_direct_booking_cutoff()
    maybe_emit_time_events()
    if request.path.startswith("/api/v1/"):
        return None  # autentificare prin token în integration_api; fără sesiune și fără CSRF
    # Authentication is valid only for the local calendar day on which the user logged in.
    auth_date = session.get("auth_date")
    if session.get("user_id") and auth_date != clock.today().isoformat():
        session.clear()
    g.user = current_user()
    endpoint = request.endpoint or ""
    if g.user:
        moment, last = time.time(), session.get("last_active")
        if last is not None and moment - float(last) > idle_limit_seconds(g.user):
            session.clear()
            g.user = None
            flash("Sesiunea a expirat din cauza inactivității. Autentificați-vă din nou.", "error")
        elif endpoint not in PASSIVE_ENDPOINTS or last is None:
            session["last_active"] = moment
    # Housekeeper accounts only see housekeeping tasks and can sign out.
    if g.user and g.user["staff_type"] == "housekeeping" and endpoint not in {"housekeeping", "logout", "static", "notifications_feed", "housekeeping_updates", "deep_clean_form", "deep_clean_save"}:
        return redirect(url_for("housekeeping"))
    if g.user and g.user["staff_type"] == "shift_lead" and endpoint not in {"yesterday_report", "logout", "static"}:
        return redirect(url_for("yesterday_report"))
    if g.user and g.user["staff_type"] == "night_reception" and endpoint not in {"night_desk", "night_checkin", "night_arrived", "night_cancel", "logout", "static"}:
        return redirect(url_for("night_desk"))
    if g.user and g.user["staff_type"] == "housekeeping" and endpoint != "static":
        track_housekeeper_presence()
    # Refuz implicit: orice rută în afara celor publice cere autentificare (nu doar cele din listă).
    if endpoint and endpoint not in PUBLIC_ENDPOINTS and not g.user:
        return redirect(url_for("setup" if not users_exist() else "login", next=request.path))
    if request.method == "POST":
        token = session.get("csrf_token")
        submitted = request.form.get("csrf_token", "")
        if not token or not secrets.compare_digest(token, submitted):
            abort(400, description="Token de securitate invalid. Reîncărcați pagina și încercați din nou.")

@app.context_processor
def inject_globals():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(32)
    return {"csrf_token": session["csrf_token"], "csp_nonce": getattr(g, "csp_nonce", ""), "current_user": getattr(g, "user", None), "today_iso": clock.today().isoformat(), "app_version": app_version_info.label(), "mods": settings.modules_state(), "times": {k: settings.time_text(k) for k, *_ in settings.TIMES}, "test_clock": clock.now().strftime("%d.%m.%Y %H:%M") if clock.is_test() else None}

def users_exist():
    conn = connect()
    exists = conn.execute("SELECT 1 FROM users LIMIT 1").fetchone() is not None
    conn.close()
    return exists

pin_valid = safety.pin_valid
pin_message = safety.pin_message

@app.route("/setup", methods=["GET", "POST"])
def setup():
    if users_exist():
        return redirect(url_for("login"))
    if request.method == "POST":
        username = request.form.get("username", "admin").strip().lower() or "admin"
        display_name = request.form.get("display_name", "Administrator").strip() or "Administrator"
        pin = request.form.get("pin", "")
        pin2 = request.form.get("pin_confirm", "")
        if not username.replace("_", "").replace("-", "").isalnum():
            flash("Utilizați pentru utilizator doar litere, cifre, _ sau -.", "error")
        elif not pin_valid(pin):
            flash(pin_message(pin), "error")
        elif pin != pin2:
            flash("PIN-urile nu coincid.", "error")
        else:
            conn = connect()
            conn.execute("INSERT INTO users(username, display_name, pin_hash, role) VALUES (?, ?, ?, 'admin')", (username, display_name, generate_password_hash(pin)))
            user = conn.execute("SELECT id FROM users WHERE username=?", (username,)).fetchone()
            conn.close()
            session.clear()
            session["user_id"] = user["id"]
            session["csrf_token"] = secrets.token_urlsafe(32)
            session.permanent = True
            g.user = current_user()
            audit("create", "user", user["id"], {"role": "admin", "username": username, "initial_setup": True})
            flash("Contul de administrator a fost creat. Adăugați separat utilizatorii recepției.", "success")
            return redirect(url_for("index"))
    return render_template("setup.html")

@app.route("/login", methods=["GET", "POST"])
def login():
    if g.user:
        return redirect(url_for("index"))
    if not users_exist():
        return redirect(url_for("setup"))
    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        pin = request.form.get("pin", "")
        username = username[:64]
        ip = request.remote_addr or "unknown"
        conn = connect()
        recent = conn.execute("SELECT COUNT(*) FROM auth_attempts WHERE ip_address=? AND success=0 AND attempted_at >= datetime('now','-15 minutes')", (ip,)).fetchone()[0]
        recent_user = conn.execute("SELECT COUNT(*) FROM auth_attempts WHERE username=? AND success=0 AND attempted_at >= datetime('now','-15 minutes')", (username,)).fetchone()[0]
        if recent >= 5 or recent_user >= 10:
            conn.close()
            flash("Prea multe încercări. Așteptați 15 minute înainte de a încerca din nou.", "error")
            return render_template("login.html")
        user = conn.execute("SELECT * FROM users WHERE username=? AND active=1", (username,)).fetchone()
        hash_to_check = user["pin_hash"] if user else DUMMY_PIN_HASH      # același cost de calcul, indiferent dacă utilizatorul există
        ok = check_password_hash(hash_to_check, pin) and bool(user)
        conn.execute("INSERT INTO auth_attempts(ip_address, success, username) VALUES (?, ?, ?)", (ip, 1 if ok else 0, username))
        if ok:
            conn.execute("UPDATE users SET last_login_at=CURRENT_TIMESTAMP WHERE id=?", (user["id"],))
        conn.close()
        if not ok:
            audit("login_failed", "user", None, {"attempted_username": username})
            flash("Utilizator sau PIN incorect.", "error")
            return render_template("login.html")
        target = safety.safe_next(request.args.get("next", ""), url_for("index"))
        if user["role"] == "admin" and (user["totp_enabled"] or settings.enabled("admin_2fa_required")):
            session.clear()                                   # PIN corect, dar contul NU este încă autentificat
            session.update(pending_user=user["id"], pending_at=time.time(), pending_next=target, csrf_token=secrets.token_urlsafe(32))
            return redirect(url_for("login_2fa" if user["totp_enabled"] else "login_2fa_setup"))
        return complete_login(user, target)
    return render_template("login.html")

def complete_login(user, target):
    session.clear()
    session["user_id"] = user["id"]
    session["auth_date"] = clock.today().isoformat()
    session["epoch"] = user["session_epoch"]
    session["last_active"] = time.time()
    session["dc_prompt"] = True  # formularul săptămânal se oferă cameristelor la autentificare, dacă este cazul
    session["csrf_token"] = secrets.token_urlsafe(32)
    session.permanent = True
    g.user = current_user()
    audit("login", "user", user["id"], {"username": user["username"], "second_factor": bool(user["totp_enabled"]) if "totp_enabled" in user.keys() else False})
    return redirect(target)


def pending_user():
    """Utilizatorul care a introdus PIN-ul corect și așteaptă al doilea pas (valabil 5 minute)."""
    if not session.get("pending_user") or time.time() - float(session.get("pending_at", 0)) > 300:
        return None
    conn = connect()
    user = conn.execute("SELECT * FROM users WHERE id=? AND active=1", (session["pending_user"],)).fetchone()
    conn.close()
    return user


def second_factor_locked(username):
    conn = connect()
    failures = conn.execute("SELECT COUNT(*) FROM auth_attempts WHERE username=? AND success=0 AND attempted_at >= datetime('now','-15 minutes')", (f"2fa:{username}",)).fetchone()[0]
    conn.close()
    return failures >= 5


@app.route("/login/2fa", methods=["GET", "POST"])
def login_2fa():
    user = pending_user()
    if not user or not user["totp_enabled"]:
        session.clear()
        flash("Sesiunea de autentificare a expirat. Introduceți din nou utilizatorul și PIN-ul.", "error")
        return redirect(url_for("login"))
    if request.method == "POST":
        if second_factor_locked(user["username"]):
            flash("Prea multe coduri greșite. Așteptați 15 minute.", "error")
            return render_template("login_2fa.html", mode="code")
        code = request.form.get("code", "").strip()
        step = totp.verify(user["totp_secret"], code, time.time(), user["totp_last_step"])   # ora reală: telefonul folosește ora reală
        conn = connect()
        used_recovery = False
        if step is None and "-" in code:
            hashes = json.loads(user["totp_recovery"] or "[]")
            match = next((h for h in hashes if check_password_hash(h, code)), None)
            if match:
                hashes.remove(match)
                conn.execute("UPDATE users SET totp_recovery=? WHERE id=?", (json.dumps(hashes), user["id"]))
                used_recovery = len(hashes)
        conn.execute("INSERT INTO auth_attempts(ip_address, success, username) VALUES (?, ?, ?)",
                     (request.remote_addr or "unknown", 1 if (step is not None or used_recovery is not False) else 0, f"2fa:{user['username']}"))
        if step is not None:
            conn.execute("UPDATE users SET totp_last_step=? WHERE id=?", (step, user["id"]))
        conn.close()
        if step is not None or used_recovery is not False:
            if used_recovery is not False:
                flash(f"Ați folosit un cod de rezervă. Mai aveți {used_recovery}. Dacă ați pierdut telefonul, cereți resetarea autentificării în 2 pași.", "error")
            return complete_login(user, session.get("pending_next") or url_for("index"))
        audit("login_2fa_failed", "user", user["id"], {})
        flash("Cod incorect sau expirat.", "error")
    return render_template("login_2fa.html", mode="code")


@app.route("/login/2fa/setup", methods=["GET", "POST"])
def login_2fa_setup():
    user = pending_user()
    if not user or user["totp_enabled"]:
        session.clear()
        return redirect(url_for("login"))
    secret = session.get("totp_setup") or totp.new_secret()
    session["totp_setup"] = secret
    if request.method == "POST":
        step = totp.verify(secret, request.form.get("code", ""), time.time())
        if step is None:
            flash("Cod incorect. Verificați că ora telefonului este corectă și introduceți codul afișat acum.", "error")
        else:
            codes = totp.new_recovery_codes()
            conn = connect()
            conn.execute("UPDATE users SET totp_secret=?, totp_enabled=1, totp_last_step=?, totp_recovery=? WHERE id=?",
                         (secret, step, json.dumps([generate_password_hash(c) for c in codes]), user["id"]))
            conn.close()
            audit("totp_enabled", "user", user["id"], {})
            target = session.get("pending_next") or url_for("index")
            fresh = dict(user)
            fresh["totp_enabled"] = 1
            response = complete_login(fresh, target)
            session["show_recovery"] = codes
            return redirect(url_for("recovery_codes"))
    return render_template("login_2fa.html", mode="setup", secret=totp.grouped(secret), uri=totp.provisioning_uri(secret, user["username"]))


@app.route("/account/recovery-codes")
@login_required
def recovery_codes():
    codes = session.pop("show_recovery", None)
    if not codes:
        return redirect(url_for("index"))
    return render_template("login_2fa.html", mode="codes", codes=codes)


@app.route("/logout", methods=["POST"])
@login_required
def logout():
    audit("logout", "user", g.user["id"])
    if session.get("ts_id"):
        conn = connect()
        timesheet.end_session(conn, session.get("ts_id"))
        conn.close()
    session.clear()
    flash("Ați ieșit din aplicație.", "success")
    return redirect(url_for("login"))

@app.route("/api/notifications")
@login_required
def notifications_feed():
    module_required("mod_alerts")
    raw = request.args.get("after_id", "")
    try:
        after_id = int(raw) if raw != "" else None
    except ValueError:
        abort(400)
    conn = connect()
    data = notifications.feed(conn, notifications.audience_for(g.user["staff_type"]), g.user["id"], after_id)
    conn.close()
    response = jsonify(data)
    response.headers["Cache-Control"] = "no-store"
    return response


def module_required(key):
    """Modul oprit de administrator = funcția nu mai răspunde (nu doar ascunsă)."""
    if not settings.enabled(key):
        abort(404)


def module_off_page(title):
    return render_template("error.html", code=403, message=f"„{title}” este dezactivat de administrator."), 403


def night_access_required():
    if g.user["role"] != "admin" and g.user["staff_type"] != "night_reception":
        abort(403)


def night_state():
    """(ora curentă, data nopții, recepția de noapte este activă?)"""
    now = room_state.now_fn()
    return now, night.night_date(now), night.is_night(now)


def night_actions_allowed():
    """Acțiunile se pot face numai în programul de noapte; administratorul este exceptat (verificări/teste)."""
    _, _, active = night_state()
    return active or g.user["role"] == "admin" or not settings.enabled("night_time_restriction")


def night_blocked_redirect():
    start, end = night.window()
    flash(f"Recepția de noapte este activă între {start.strftime('%H:%M')} și {end.strftime('%H:%M')}. Acum puteți doar vizualiza.", "error")
    return redirect(url_for("night_desk"))


@app.route("/night")
@login_required
def night_desk():
    night_access_required()
    if not settings.enabled("mod_night_desk"):
        return module_off_page("Recepția de noapte")
    now, nd, active = night_state()
    conn = connect()
    rooms = night.grid(conn, nd)
    late = night.late_arrivals(conn, nd)
    conn.close()
    buildings = []
    for name in ("Corp A", "Corp B", "Corp C"):
        group = [r for r in rooms if r["building"] == name]
        if group:
            buildings.append({"name": name, "rooms": group})
    counts = {}
    for r in rooms:
        counts[r["state"]] = counts.get(r["state"], 0) + 1
    start, end = night.window()
    return render_template("night.html", night=nd, next_day=nd + timedelta(days=1), buildings=buildings, late=late, counts=counts,
                           can_act=night_actions_allowed(), start=start.strftime("%H:%M"), end=end.strftime("%H:%M"),
                           payments=night.PAYMENTS)


@app.route("/night/room/<int:room_id>", methods=["POST"])
@login_required
def night_checkin(room_id):
    night_access_required()
    module_required("mod_night_desk")
    if not night_actions_allowed():
        return night_blocked_redirect()
    _, nd, _ = night_state()
    mode = request.form.get("mode", "")
    conn = connect()
    ok, message, booking_id = night.create_for_free_room(conn, room_id, nd, mode, request.form.get("guest_name"), request.form.get("guests"),
                                                         request.form.get("payment"), g.user, color_for_user(g.user["id"]))
    conn.close()
    if ok:
        audit("night_checkin" if mode == "occupy" else "night_reserve", "booking", booking_id,
              {"room_id": room_id, "status": "checked_in" if mode == "occupy" else "confirmed", "payment": request.form.get("payment") if mode == "occupy" else None})
    flash(message, "success" if ok else "error")
    return redirect(url_for("night_desk"))


@app.route("/night/booking/<int:booking_id>/arrived", methods=["POST"])
@login_required
def night_arrived(booking_id):
    night_access_required()
    module_required("mod_night_desk")
    if not night_actions_allowed():
        return night_blocked_redirect()
    _, nd, _ = night_state()
    conn = connect()
    ok, message, row = night.mark_arrived(conn, booking_id, nd, request.form.get("payment"), g.user)
    conn.close()
    if ok:
        audit("night_arrived", "booking", booking_id, {"room_id": row["room_id"], "status": "checked_in", "payment": request.form.get("payment")})
    flash(message, "success" if ok else "error")
    return redirect(url_for("night_desk"))


@app.route("/night/booking/<int:booking_id>/cancel", methods=["POST"])
@login_required
def night_cancel(booking_id):
    night_access_required()
    module_required("mod_night_desk")
    if not night_actions_allowed():
        return night_blocked_redirect()
    _, nd, _ = night_state()
    conn = connect()
    ok, message, row = night.cancel_late(conn, booking_id, nd, g.user)
    conn.close()
    if ok:
        audit("night_cancel", "booking", booking_id, {"room_id": row["room_id"], "status": "cancelled"})
    flash(message, "success" if ok else "error")
    return redirect(url_for("night_desk"))


@app.route("/booking/<int:booking_id>/night-payment-collected", methods=["POST"])
@login_required
def night_payment_collected(booking_id):
    """Recepția de zi marchează încasată plata cazării făcute noaptea."""
    module_required("mod_night_pending")
    conn = connect()
    row = conn.execute("SELECT id, night_payment FROM bookings WHERE id=?", (booking_id,)).fetchone()
    if not row or row["night_payment"] != "morning":
        conn.close()
        abort(404)
    conn.execute("UPDATE bookings SET night_payment='collected', updated_at=CURRENT_TIMESTAMP WHERE id=?", (booking_id,))
    conn.close()
    audit("night_payment_collected", "booking", booking_id, {})
    flash("Plata a fost marcată ca încasată.", "success")
    return redirect(url_for("index"))


@app.route("/admin/rooms/rate", methods=["POST"])
@admin_required
def set_room_rate():
    try:
        room_id = int(request.form.get("room_id", "0"))
        rate = round(float(request.form.get("default_rate", "")), 2)
    except ValueError:
        abort(400)
    if not 0 <= rate <= 10000:
        flash("Tariful trebuie să fie între 0 și 10.000 lei.", "error")
        return redirect(url_for("room_maintenance"))
    conn = connect()
    room = conn.execute("SELECT id, name FROM rooms WHERE id=? AND active=1", (room_id,)).fetchone()
    if not room:
        conn.close()
        abort(404)
    conn.execute("UPDATE rooms SET default_rate=? WHERE id=?", (rate, room_id))
    conn.close()
    audit("room_rate_set", "room", room_id, {"room": room["name"], "default_rate": rate})
    flash(f"Tariful camerei {room['name']} a fost salvat: {rate:.2f} lei / noapte.", "success")
    return redirect(url_for("room_maintenance"))


@app.route("/report/yesterday")
@login_required
def yesterday_report():
    """Raport simplu de ocupare: ce camere au fost ocupate noaptea precedentă (pentru șeful de tură)."""
    if not settings.enabled("mod_shift_report"):
        return module_off_page("Raportul șefului de tură")
    night = clock.today() - timedelta(days=1)
    if g.user["staff_type"] != "shift_lead":
        try:
            night = datetime.strptime(request.args.get("date", ""), "%Y-%m-%d").date()
        except ValueError:
            pass
    conn = connect()
    rooms = conn.execute("SELECT id, name, building FROM rooms WHERE active=1 ORDER BY CASE building WHEN 'Corp A' THEN 1 WHEN 'Corp B' THEN 2 WHEN 'Corp C' THEN 3 ELSE 4 END, sort_order, name").fetchall()
    stays = conn.execute("""SELECT room_id, rate, check_in, night_payment FROM bookings
        WHERE status IN ('confirmed','checked_in','checked_out') AND check_in<=? AND check_out>?""", (night.isoformat(), night.isoformat())).fetchall()
    conn.close()
    occupied_ids = {r["room_id"] for r in stays}
    active_ids = {r["id"] for r in rooms}
    # Total încasări pe noapte = tariful pe noapte al fiecărei camere ocupate (aplicația nu ține evidența plăților separat).
    revenue = round(sum(float(r["rate"] or 0) for r in stays if r["room_id"] in active_ids), 2)
    night_desk = [r for r in stays if r["night_payment"] and r["check_in"] == night.isoformat() and r["room_id"] in active_ids]
    cash = round(sum(float(r["rate"] or 0) for r in night_desk if r["night_payment"] == "cash"), 2)
    morning = round(sum(float(r["rate"] or 0) for r in night_desk if r["night_payment"] in ("morning", "collected")), 2)
    buildings = []
    for name in ("Corp A", "Corp B", "Corp C"):
        names = [r["name"] for r in rooms if r["building"] == name and r["id"] in occupied_ids]
        total = sum(1 for r in rooms if r["building"] == name)
        if total:
            buildings.append({"name": name, "occupied": names, "total": total})
    total_rooms = len(rooms)
    occupied = len(occupied_ids & {r["id"] for r in rooms})
    percent = round(100 * occupied / total_rooms) if total_rooms else 0
    return render_template("yesterday_report.html", night=night, next_day=night + timedelta(days=1), buildings=buildings,
                           occupied=occupied, total=total_rooms, percent=percent, revenue=revenue,
                           night_desk_count=len(night_desk), night_cash=cash, night_morning=morning)


@app.route("/admin/timesheet")
@admin_required
def timesheet_page():
    module_required("mod_timesheet")
    try:
        day = datetime.strptime(request.args.get("date", ""), "%Y-%m-%d").date()
    except ValueError:
        day = clock.today()
    monday = day - timedelta(days=day.weekday())
    conn = connect()
    rows = timesheet.summarize_day(conn, day.isoformat())
    week_days, week_rows = timesheet.week_totals(conn, monday)
    conn.close()
    return render_template("timesheet.html", day=day, rows=rows, week_days=week_days, week_rows=week_rows,
                           prev_day=(day - timedelta(days=1)).isoformat(), next_day=(day + timedelta(days=1)).isoformat(),
                           is_today=(day == clock.today()))


@app.route("/admin/timesheet.csv")
@admin_required
def timesheet_csv():
    module_required("mod_timesheet")
    try:
        year, month = int(request.args.get("year", clock.today().year)), int(request.args.get("month", clock.today().month))
        if not 1 <= month <= 12:
            raise ValueError
    except ValueError:
        abort(400)
    conn = connect()
    body = timesheet.month_csv(conn, year, month)
    conn.close()
    response = Response(body, mimetype="text/csv")
    response.headers["Content-Disposition"] = f"attachment; filename=pontaj_{year:04d}-{month:02d}.csv"
    return response


@app.route("/admin/settings", methods=["GET", "POST"])
@admin_required
def admin_settings():
    if request.method == "POST":
        values = {k: ("1" if request.form.get(k) == "1" else "0") for k, *_ in settings.MODULES}
        values.update({k: request.form.get(k, "").strip() for k, *_ in settings.TIMES})
        values.update({k: request.form.get(k, "").strip() for k, *_ in settings.NUMBERS})
        conn = connect()
        ok, message, changed = settings.save(conn, values, g.user["username"])
        conn.close()
        if ok and changed:
            audit("settings_update", "settings", None, {k: values[k] for k in changed})
        flash(message if changed or not ok else "Nicio modificare.", "success" if ok else "error")
        return redirect(url_for("admin_settings"))
    return render_template("settings.html", modules=settings.MODULES, state=settings.modules_state(),
                           time_rows=[(k, title, settings.time_text(k)) for k, title, *_ in settings.TIMES],
                           number_rows=[(k, title, settings.number(k), lo, hi) for k, title, _, lo, hi in settings.NUMBERS])


@app.route("/admin/backups", methods=["GET", "POST"])
@admin_required
def admin_backups():
    dest = backup.folder(DB_PATH)
    if request.method == "POST":
        try:
            result = backup.run(DB_PATH, dest, settings.number("backup_keep_days"), clock.now())
            audit("backup_manual", "backup", None, {"file": os.path.basename(result["path"]), "bookings": result["bookings"]})
            flash(f"Copie salvată și verificată: {os.path.basename(result['path'])} ({result['bookings']} rezervări).", "success")
        except Exception:
            app.logger.exception("Copia de siguranță a eșuat")
            flash("Copia de siguranță a eșuat. Verificați spațiul liber și drepturile asupra folderului.", "error")
        return redirect(url_for("admin_backups"))
    return render_template("backups.html", items=backup.list_backups(dest), folder=dest, keep=settings.number("backup_keep_days"),
                           at=settings.time_text("backup_time"))


ROLE_CHOICES = [("admin", "Administrator"), ("reception", "Recepție"), ("housekeeping", "Cameristă"), ("shift_lead", "Șef de tură"), ("night_reception", "Recepție noapte")]


@app.route("/admin/users/<int:user_id>/edit", methods=["GET", "POST"])
@admin_required
def edit_user(user_id):
    conn = connect()
    target = conn.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
    if not target:
        conn.close()
        abort(404)
    current_role = target["role"] if target["role"] == "admin" else (target["staff_type"] or "reception")
    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        display_name = request.form.get("display_name", "").strip()
        role = request.form.get("role", "")
        pin = request.form.get("pin", "")
        active = 1 if request.form.get("active") == "1" else 0
        reset_2fa = request.form.get("reset_2fa") == "1"
        logout_all = request.form.get("logout_all") == "1"
        is_self = target["id"] == g.user["id"]
        other_admins = conn.execute("SELECT COUNT(*) FROM users WHERE role='admin' AND active=1 AND id!=?", (user_id,)).fetchone()[0]
        error = None
        if not username.replace("_", "").replace("-", "").isalnum() or not display_name or len(username) > 32 or len(display_name) > 60:
            error = "Completați un utilizator valid (litere, cifre, - și _) și numele afișat."
        elif role not in dict(ROLE_CHOICES):
            error = "Rol invalid."
        elif pin and not pin_valid(pin):
            error = pin_message(pin)
        elif is_self and (role != "admin" or not active):
            error = "Nu vă puteți retrage propriul rol de administrator și nu vă puteți dezactiva propriul cont."
        elif target["role"] == "admin" and (role != "admin" or not active) and other_admins == 0:
            error = "Trebuie să rămână cel puțin un administrator activ."
        if error:
            conn.close()
            flash(error, "error")
            return render_template("user_edit.html", u=target, current_role=current_role, roles=ROLE_CHOICES, form=request.form)
        staff_type = role if role in ("housekeeping", "shift_lead", "night_reception") else "reception"
        db_role = "admin" if role == "admin" else "reception"
        changed = [name for name, old, new in (("username", target["username"], username), ("display_name", target["display_name"], display_name),
                   ("role", current_role, role), ("active", target["active"], active)) if old != new]
        if pin:
            changed.append("pin")
        if reset_2fa and target["totp_enabled"]:
            changed.append("2fa_reset")
        if logout_all:
            changed.append("logout_all")
        invalidate = any(c in changed for c in ("role", "active", "pin", "2fa_reset", "logout_all"))
        try:
            conn.execute("UPDATE users SET username=?, display_name=?, role=?, staff_type=?, active=?, session_epoch=session_epoch+? WHERE id=?",
                         (username, display_name, db_role, staff_type, active, 1 if invalidate else 0, user_id))
        except sqlite3.IntegrityError:
            conn.close()
            flash("Acest nume de utilizator există deja.", "error")
            return render_template("user_edit.html", u=target, current_role=current_role, roles=ROLE_CHOICES, form=request.form)
        if pin:
            conn.execute("UPDATE users SET pin_hash=? WHERE id=?", (generate_password_hash(pin), user_id))
        if "2fa_reset" in changed:
            conn.execute("UPDATE users SET totp_secret=NULL, totp_enabled=0, totp_last_step=NULL, totp_recovery='[]' WHERE id=?", (user_id,))
        if is_self and invalidate:
            session["epoch"] = conn.execute("SELECT session_epoch FROM users WHERE id=?", (user_id,)).fetchone()[0]   # nu vă deconectați pe dumneavoastră
        conn.close()
        if changed:
            audit("update", "user", user_id, {"fields": changed, "username": username})
        flash("Utilizatorul a fost actualizat." if changed else "Nicio modificare.", "success")
        return redirect(url_for("users"))
    conn.close()
    return render_template("user_edit.html", u=target, current_role=current_role, roles=ROLE_CHOICES, form=None)


@app.route("/admin/users", methods=["GET", "POST"])
@admin_required
def users():
    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        display_name = request.form.get("display_name", "").strip()
        pin = request.form.get("pin", "")
        role = request.form.get("role", "reception")
        staff_type = role if role in ("housekeeping", "shift_lead", "night_reception") else "reception"
        db_role = "reception" if role in ("housekeeping", "shift_lead", "night_reception") else role
        if not username.replace("_", "").replace("-", "").isalnum() or not display_name or len(username) > 32 or len(display_name) > 60:
            flash("Completați un utilizator valid și numele afișat.", "error")
        elif not pin_valid(pin):
            flash(pin_message(pin), "error")
        elif role not in {"admin", "reception", "housekeeping", "shift_lead", "night_reception"}:
            flash("Rol invalid.", "error")
        else:
            conn = connect()
            try:
                cur = conn.execute("INSERT INTO users(username, display_name, pin_hash, role, staff_type) VALUES (?, ?, ?, ?, ?)", (username, display_name, generate_password_hash(pin), db_role, staff_type))
                conn.close()
                audit("create", "user", cur.lastrowid, {"username": username, "role": role})
                flash("Utilizatorul a fost creat.", "success")
                return redirect(url_for("users"))
            except sqlite3.IntegrityError:
                conn.close()
                flash("Acest nume de utilizator există deja.", "error")
    conn = connect()
    all_users = conn.execute("SELECT id, username, display_name, role, COALESCE(staff_type,'reception') AS staff_type, active, created_at, last_login_at FROM users ORDER BY id").fetchall()
    conn.close()
    return render_template("users.html", users=all_users, user_colors=USER_COLORS)

@app.route("/admin/users/<int:user_id>/toggle", methods=["POST"])
@admin_required
def toggle_user(user_id):
    if user_id == g.user["id"]:
        flash("Nu vă puteți dezactiva propriul cont.", "error")
        return redirect(url_for("users"))
    conn = connect()
    user = conn.execute("SELECT active, username FROM users WHERE id=?", (user_id,)).fetchone()
    if not user:
        conn.close()
        abort(404)
    active = 0 if user["active"] else 1
    conn.execute("UPDATE users SET active=? WHERE id=?", (active, user_id))
    conn.close()
    audit("activate" if active else "deactivate", "user", user_id, {"username": user["username"]})
    flash("Starea utilizatorului a fost actualizată.", "success")
    return redirect(url_for("users"))

@app.route("/admin/users/<int:user_id>/delete", methods=["POST"])
@admin_required
def delete_user(user_id):
    if user_id == g.user["id"]:
        flash("Nu vă puteți șterge propriul cont.", "error")
        return redirect(url_for("users"))
    conn = connect()
    user = conn.execute("SELECT id, username, display_name, role FROM users WHERE id=?", (user_id,)).fetchone()
    if not user:
        conn.close()
        abort(404)
    if user["role"] == "admin":
        admin_count = conn.execute("SELECT COUNT(*) FROM users WHERE role='admin'").fetchone()[0]
        if admin_count <= 1:
            conn.close()
            flash("Nu poate fi șters ultimul cont de administrator.", "error")
            return redirect(url_for("users"))
    # Preserve historical audit/invoice records; their foreign keys are SET NULL.
    deleted_identity = {"username": user["username"], "display_name": user["display_name"], "role": user["role"]}
    conn.execute("DELETE FROM users WHERE id=?", (user_id,))
    conn.close()
    audit("delete", "user", user_id, deleted_identity)
    flash(f"Utilizatorul {user['username']} a fost șters. Istoricul acțiunilor sale rămâne în audit.", "success")
    return redirect(url_for("users"))

@app.route("/admin/users/<int:user_id>/pin", methods=["POST"])
@admin_required
def reset_user_pin(user_id):
    pin = request.form.get("pin", "")
    if not pin_valid(pin):
        flash(pin_message(pin), "error")
        return redirect(url_for("users"))
    conn = connect()
    user = conn.execute("SELECT username FROM users WHERE id=?", (user_id,)).fetchone()
    if not user:
        conn.close(); abort(404)
    conn.execute("UPDATE users SET pin_hash=?, session_epoch=session_epoch+1 WHERE id=?", (generate_password_hash(pin), user_id))
    conn.close()
    audit("reset_pin", "user", user_id, {"username": user["username"]})
    flash("PIN-ul a fost resetat. Comunicați-l angajatului printr-un canal separat.", "success")
    return redirect(url_for("users"))

@app.route("/admin/audit")
@admin_required
def audit_page():
    audit("view", "audit_log", None, {"limit": 200})
    conn = connect()
    entries = conn.execute("SELECT * FROM audit_log ORDER BY id DESC LIMIT 200").fetchall()
    conn.close()
    return render_template("audit.html", entries=entries)

@app.route("/admin/audit/export")
@admin_required
def audit_export():
    audit("export", "audit_log", None, {"format": "txt"})
    conn = connect()
    entries = conn.execute("SELECT created_at, username, user_id, ip_address, action, entity_type, entity_id, details FROM audit_log ORDER BY id ASC").fetchall()
    conn.close()
    lines = ["POPAS CĂPRIOARA PMS - JURNAL DE AUDIT", "Export generat: " + clock.now().astimezone().isoformat(timespec="seconds"), "=" * 100]
    for e in entries:
        lines.append(f"[{e['created_at']}] utilizator={e['username']} (id={e['user_id'] or '-'}) | IP={e['ip_address'] or '-'} | acțiune={e['action']} | obiect={e['entity_type']}#{e['entity_id'] or '-'} | detalii={e['details']}")
    body = "\r\n".join(lines) + "\r\n"
    from flask import Response
    return Response(body, mimetype="text/plain; charset=utf-8", headers={"Content-Disposition": "attachment; filename=popas_caprioara_audit.txt"})



@app.route("/admin/rooms", methods=["GET", "POST"])
@admin_required
def room_maintenance():
    conn = connect()
    if request.method == "POST":
        try:
            room_id = int(request.form.get("room_id", "0"))
        except ValueError:
            conn.close(); abort(400)
        out_of_order = 1 if request.form.get("out_of_order") == "1" else 0
        note = request.form.get("maintenance_note", "").strip()[:500]
        room = conn.execute("SELECT id, name, out_of_order FROM rooms WHERE id=? AND active=1", (room_id,)).fetchone()
        if not room:
            conn.close(); abort(404)
        future_bookings = conn.execute("SELECT COUNT(*) FROM bookings WHERE room_id=? AND check_out>? AND status NOT IN ('cancelled','checked_out')", (room_id, clock.today().isoformat())).fetchone()[0] if out_of_order else 0
        conn.execute("UPDATE rooms SET out_of_order=?, maintenance_note=? WHERE id=?", (out_of_order, note, room_id))
        conn.close()
        audit("room_out_of_order" if out_of_order else "room_returned_to_service", "room", room_id, {"room": room["name"], "maintenance_note": note, "existing_future_bookings": future_bookings})
        message = f"Camera {room['name']} a fost marcată ca defectă."
        if out_of_order and future_bookings:
            message += f" Atenție: există {future_bookings} rezervare(i) viitoare pe această cameră; verificați și realocați-le manual."
        if not out_of_order:
            message = f"Camera {room['name']} a fost reactivată."
        flash(message, "success" if not future_bookings else "error")
        return redirect(url_for("room_maintenance"))
    rooms = conn.execute("SELECT * FROM rooms WHERE active=1 ORDER BY building, sort_order, name").fetchall()
    conn.close()
    return render_template("rooms.html", rooms=rooms)

@app.route("/api/available-rooms")
@login_required
def available_rooms():
    """Return only rooms that can actually be booked for the selected stay."""
    start_raw = request.args.get("start", "")
    end_raw = request.args.get("end", "")
    try:
        start = date.fromisoformat(start_raw)
        end = date.fromisoformat(end_raw)
    except (ValueError, TypeError):
        return jsonify({"rooms": [], "error": "Introduceți date valide."}), 400
    if end <= start:
        return jsonify({"rooms": [], "error": "Data plecării trebuie să fie după sosire."}), 400
    exclude_id = request.args.get("exclude_booking_id", type=int)
    conn = connect()
    # Editing a reservation must not count that same reservation as a conflict.
    # The exclusion is accepted only for an existing booking ID.
    if exclude_id:
        exists = conn.execute("SELECT id FROM bookings WHERE id=?", (exclude_id,)).fetchone()
        if not exists:
            exclude_id = None
    sql = """
        SELECT r.id, r.name, r.building, r.bed_configuration, r.bathroom_type
        FROM rooms r WHERE r.active=1 AND r.out_of_order=0
          AND NOT EXISTS (
            SELECT 1 FROM bookings b WHERE b.room_id=r.id
              AND b.status NOT IN ('cancelled','unconfirmed')
              AND b.check_in < ? AND b.check_out > ?
              AND (? IS NULL OR b.id != ?)
          )
        ORDER BY CASE r.building WHEN 'Corp A' THEN 1 WHEN 'Corp B' THEN 2 WHEN 'Corp C' THEN 3 ELSE 4 END,
                 r.sort_order, r.name
    """
    rows = conn.execute(sql, (end.isoformat(), start.isoformat(), exclude_id, exclude_id)).fetchall()
    conn.close()
    return jsonify({"rooms": [dict(row) for row in rows]})

@app.route("/api/guest-name-suggestions")
@login_required
def guest_name_suggestions():
    """Return recent distinct guest/group names to power local autocomplete."""
    query = request.args.get("q", "").strip()[:80]
    if len(query) < 2:
        return jsonify({"names": []})
    conn = connect()
    rows = conn.execute(
        """SELECT guest_name, MAX(id) AS latest_id FROM bookings
           WHERE guest_name LIKE ? COLLATE NOCASE
           GROUP BY guest_name ORDER BY latest_id DESC LIMIT 12""",
        (f"%{query}%",)
    ).fetchall()
    conn.close()
    return jsonify({"names": [row["guest_name"] for row in rows]})

@app.route("/housekeeping", methods=["GET", "POST"])
@login_required
def housekeeping():
    today = clock.today()
    conn = connect()
    if request.method == "POST":
        action = request.form.get("action", "mark_clean")
        try:
            if action == "complete_request":
                if g.user["staff_type"] != "housekeeping" and g.user["role"] != "admin":
                    conn.close()
                    abort(403)
                request_id = int(request.form.get("request_id", "0"))
                task = conn.execute("SELECT * FROM housekeeping_requests WHERE id=? AND service_date=? AND status='open'", (request_id, today.isoformat())).fetchone()
                if not task:
                    conn.close()
                    abort(404, description="Solicitarea nu mai este activă sau aparține altei zile.")
                conn.execute("UPDATE housekeeping_requests SET status='done', completed_by=?, completed_by_name=?, completed_at=CURRENT_TIMESTAMP WHERE id=?",
                             (g.user["id"], g.user["display_name"], request_id))
                conn.commit()
                audit("housekeeping_request_done", "housekeeping_request", request_id, {"room_id": task["room_id"], "type": task["request_type"]})
                flash("Solicitarea a fost marcată ca rezolvată și recepția va vedea actualizarea.", "success")
                conn.close()
                return redirect(url_for("housekeeping"))
            if g.user["staff_type"] != "housekeeping" and g.user["role"] != "admin":
                conn.close()
                abort(403)
            room_id = int(request.form.get("room_id", "0"))
            clean = 1 if request.form.get("is_clean") == "1" else 0
            room = conn.execute("SELECT id, name, out_of_order FROM rooms WHERE id=? AND active=1", (room_id,)).fetchone()
            if not room:
                conn.close()
                abort(404)
            if room["out_of_order"] and clean:
                conn.close()
                flash("Camera este marcată defectă. Rezolvați mentenanța înainte de a o confirma curată pentru cazare.", "error")
                return redirect(url_for("housekeeping"))
            current = room_state.compute_room_states(conn, room_id=room_id)
            if clean and current and current[0]["state"] in ("occupied", "departing"):
                conn.close()
                flash("Camera este ocupată sau ora de check-out nu a trecut încă. După ora de check-out camera devine disponibilă automat și poate fi marcată curată.", "error")
                return redirect(url_for("housekeeping"))
            conn.execute("""INSERT INTO room_cleaning(room_id, service_date, is_clean, updated_by, updated_by_name, updated_at)
                VALUES (?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(room_id, service_date) DO UPDATE SET is_clean=excluded.is_clean,
                updated_by=excluded.updated_by, updated_by_name=excluded.updated_by_name, updated_at=CURRENT_TIMESTAMP""",
                (room_id, today.isoformat(), clean, g.user["id"], g.user["display_name"]))
            conn.commit()
            audit("housekeeping_clean" if clean else "housekeeping_uncheck", "room", room_id,
                  {"date": today.isoformat(), "is_clean": bool(clean)})
            conn.close()
            flash("Starea camerei a fost actualizată și este vizibilă imediat la recepție.", "success")
            return redirect(url_for("housekeeping"))
        except (ValueError, TypeError):
            conn.close()
            abort(400, description="Date invalide.")
    rooms = conn.execute("""SELECT r.*, CASE WHEN EXISTS (SELECT 1 FROM bookings b WHERE b.room_id=r.id AND b.status='checked_out' AND b.check_out=?) THEN 1 ELSE 0 END AS departed_today
        FROM rooms r WHERE r.active=1 ORDER BY CASE r.building WHEN 'Corp A' THEN 1 WHEN 'Corp B' THEN 2 WHEN 'Corp C' THEN 3 ELSE 4 END, r.sort_order, r.name""",
        (today.isoformat(),)).fetchall()
    states = {item["room_id"]: item for item in room_state.compute_room_states(conn)}
    dc_info = deep_clean.due_info(conn, list(states.values()), today)
    if session.pop("dc_prompt", None) and dc_info["due"] and g.user["staff_type"] == "housekeeping" and settings.enabled("mod_deep_clean"):
        conn.close()
        return redirect(url_for("deep_clean_form"))
    marks = conn.execute("SELECT room_id, service_date, is_clean, updated_by, updated_by_name, updated_at FROM room_cleaning WHERE service_date=?", (today.isoformat(),)).fetchall()
    tasks = conn.execute("""SELECT h.*, r.name AS room_name, r.building FROM housekeeping_requests h JOIN rooms r ON r.id=h.room_id
        WHERE h.service_date=? AND h.status IN ('open','done') ORDER BY CASE h.status WHEN 'open' THEN 0 ELSE 1 END, r.building, r.sort_order, r.name, h.id DESC""", (today.isoformat(),)).fetchall()
    conn.close()
    clean_map = {m["room_id"]: {"clean": bool(m["is_clean"]), "updated_at": m["updated_at"], "updated_by": m["updated_by"], "updated_by_name": m["updated_by_name"]} for m in marks}
    task_items = []
    for row in tasks:
        item = dict(row)
        try:
            stamp = datetime.strptime(item.get('created_at', ''), '%Y-%m-%d %H:%M:%S').replace(tzinfo=__import__('datetime').timezone.utc).astimezone()
            item['requested_time'] = stamp.strftime('%H:%M')
        except (ValueError, TypeError):
            item['requested_time'] = (item.get('created_at') or '')[-8:-3]
        task_items.append(item)
    return render_template("housekeeping.html", rooms=rooms, today=today, clean_map=clean_map, tasks=task_items, states=states, dc=dc_info)

def housekeeping_or_admin_only():
    if g.user["staff_type"] != "housekeeping" and g.user["role"] != "admin":
        abort(403)


@app.route("/housekeeping/deep-clean")
@login_required
def deep_clean_form():
    housekeeping_or_admin_only()
    module_required("mod_deep_clean")
    conn = connect()
    states = room_state.compute_room_states(conn)
    week = deep_clean.week_start(clock.today())
    reports = deep_clean.reports_for_week(conn, week)
    info = deep_clean.due_info(conn, states)
    conn.close()
    rows = [{"state": item, "reason": deep_clean.eligibility(item), "report": reports.get(item["room_id"])} for item in states]
    return render_template("deep_clean.html", rows=rows, items=deep_clean.ITEMS, info=info, week=week, note_max=deep_clean.NOTE_MAX)


@app.route("/housekeeping/deep-clean/<int:room_id>", methods=["POST"])
@login_required
def deep_clean_save(room_id):
    housekeeping_or_admin_only()
    module_required("mod_deep_clean")
    send = request.form.get("action") == "send"
    checked = set(request.form.getlist("items")) & set(deep_clean.ITEM_KEYS)
    conn = connect()
    states = room_state.compute_room_states(conn, room_id=room_id)
    if not states:
        conn.close()
        abort(404)
    reason = deep_clean.eligibility(states[0])
    if reason:
        conn.close()
        flash(f"Camera {states[0]['room_name']} nu poate fi curățată extins acum ({reason}).", "error")
        return redirect(url_for("deep_clean_form"))
    week = deep_clean.week_start(clock.today())
    result = deep_clean.save_report(conn, week, room_id, checked, request.form.get("note", ""), g.user, send)
    conn.commit()
    conn.close()
    audit("deep_clean_submit" if send else "deep_clean_save", "room", room_id,
          {"week": week, "done": result["done"], "total": result["total"], "has_note": result["has_note"]})
    flash(f"Raportul pentru camera {states[0]['room_name']} a fost trimis recepției." if send else f"Camera {states[0]['room_name']}: salvat ca ciornă.", "success")
    return redirect(url_for("deep_clean_form") + f"#room-{room_id}")


@app.route("/deep-clean/reports")
@login_required
def deep_clean_reports():
    module_required("mod_deep_clean")
    try:
        picked = datetime.strptime(request.args.get("week", ""), "%Y-%m-%d").date()
    except ValueError:
        picked = clock.today()
    week = deep_clean.week_start(picked)
    monday = date.fromisoformat(week)
    conn = connect()
    rooms = conn.execute("SELECT id, name, building FROM rooms WHERE active=1 ORDER BY CASE building WHEN 'Corp A' THEN 1 WHEN 'Corp B' THEN 2 WHEN 'Corp C' THEN 3 ELSE 4 END, sort_order, name").fetchall()
    reports = deep_clean.reports_for_week(conn, week)
    conn.close()
    sent = sum(1 for r in reports.values() if r["status"] == "sent")
    return render_template("deep_clean_reports.html", rooms=rooms, reports=reports, items=deep_clean.ITEMS, week=week, sent=sent,
                           prev_week=(monday - timedelta(days=7)).isoformat(), next_week=(monday + timedelta(days=7)).isoformat(),
                           is_current=(week == deep_clean.week_start(clock.today())))


@app.route("/housekeeping/request", methods=["POST"])
@login_required
def create_housekeeping_request():
    module_required("mod_housekeeping_comm")
    if g.user["staff_type"] == "housekeeping":
        abort(403)
    try:
        room_id = int(request.form.get("room_id", "0"))
    except ValueError:
        abort(400, description="Camera selectată nu este validă.")
    raw_services = request.form.getlist("services") or [request.form.get("request_type", "cleaning")]
    if any(item not in set(notifications.SERVICE_ORDER) | {"both"} for item in raw_services):
        abort(400, description="Tip de solicitare invalid.")
    chosen = set()
    for item in raw_services:
        chosen.update(notifications.parse_services(item))
    request_type = ",".join(part for part in notifications.SERVICE_ORDER if part in chosen)
    if not request_type:
        abort(400, description="Alegeți cel puțin un serviciu.")
    note = request.form.get("note", "").strip()[:300]
    conn = connect()
    room = conn.execute("SELECT id, name FROM rooms WHERE id=? AND active=1 AND out_of_order=0", (room_id,)).fetchone()
    if not room:
        conn.close()
        flash("Camera nu este activă sau este marcată defectă.", "error")
        return redirect(url_for("index"))
    existing = conn.execute("SELECT id FROM housekeeping_requests WHERE room_id=? AND service_date=? AND request_type=? AND status='open'", (room_id, clock.today().isoformat(), request_type)).fetchone()
    if existing:
        conn.close()
        flash(f"Există deja o solicitare deschisă pentru camera {room['name']} de acest tip.", "error")
        return redirect(url_for("index"))
    cur = conn.execute("""INSERT INTO housekeeping_requests(room_id, service_date, request_type, note, requested_by, requested_by_name)
        VALUES (?, ?, ?, ?, ?, ?)""", (room_id, clock.today().isoformat(), request_type, note, g.user["id"], g.user["display_name"]))
    request_id = cur.lastrowid
    conn.commit(); conn.close()
    audit("housekeeping_request_created", "housekeeping_request", request_id, {"room_id": room_id, "room": room["name"], "type": request_type})
    flash(f"Solicitarea pentru camera {room['name']} a fost transmisă cameristelor.", "success")
    return redirect(url_for("index"))

@app.route("/api/housekeeping-updates")
@login_required
def housekeeping_updates():
    """Small polling endpoint for near-real-time reception/housekeeping updates."""
    today_iso = clock.today().isoformat()
    conn = connect()
    open_rows = conn.execute("""SELECT h.id, h.room_id, r.name AS room_name, h.request_type, h.note,
        h.status, h.created_at, h.completed_at, h.completed_by_name
        FROM housekeeping_requests h JOIN rooms r ON r.id=h.room_id
        WHERE h.service_date=? AND h.status IN ('open','done')
        ORDER BY h.id DESC LIMIT 100""", (today_iso,)).fetchall()
    latest_booking_update = conn.execute("SELECT COALESCE(MAX(updated_at), '') FROM bookings WHERE check_in<=? AND check_out>=?", (today_iso, today_iso)).fetchone()[0]
    latest_cleaning_update = conn.execute("SELECT COALESCE(MAX(updated_at), '') FROM room_cleaning WHERE service_date=?", (today_iso,)).fetchone()[0]
    conn.close()
    return jsonify({"requests": [dict(row) for row in open_rows], "latest_booking_update": latest_booking_update, "latest_cleaning_update": latest_cleaning_update})

@app.route("/admin/reports")
@admin_required
def reports():
    start_raw = request.args.get("start", (clock.today() - timedelta(days=29)).isoformat())
    end_raw = request.args.get("end", clock.today().isoformat())
    try:
        start_date = date.fromisoformat(start_raw)
        end_date = date.fromisoformat(end_raw)
    except ValueError:
        flash("Perioada raportului conține o dată invalidă.", "error")
        start_date = clock.today() - timedelta(days=29)
        end_date = clock.today()
    if end_date < start_date:
        flash("Data finală trebuie să fie egală sau după data de început.", "error")
        start_date, end_date = end_date, start_date
    period_end_exclusive = end_date + timedelta(days=1)
    period_days = (period_end_exclusive - start_date).days
    conn = connect()
    rooms = conn.execute("SELECT id, name, building, out_of_order FROM rooms WHERE active=1 ORDER BY sort_order, name").fetchall()
    bookings = conn.execute("""
        SELECT b.*, r.name AS room_name, r.building AS building
        FROM bookings b JOIN rooms r ON r.id=b.room_id
        WHERE r.active=1 AND b.status IN ('confirmed', 'checked_in', 'checked_out') AND b.check_in < ? AND b.check_out > ?
        ORDER BY b.check_in, r.sort_order, r.name
    """, (period_end_exclusive.isoformat(), start_date.isoformat())).fetchall()
    conn.close()
    sold_nights = 0
    estimated_revenue = 0.0
    overlapping_bookings = 0
    room_stats = {r["id"]: {"room_name": r["name"], "building": r["building"], "out_of_order": bool(r["out_of_order"]), "sold_nights": 0, "revenue": 0.0, "bookings": 0} for r in rooms}
    building_stats = {}
    for b in bookings:
        clipped_start = max(date.fromisoformat(b["check_in"]), start_date)
        clipped_end = min(date.fromisoformat(b["check_out"]), period_end_exclusive)
        nights = max(0, (clipped_end - clipped_start).days)
        if not nights:
            continue
        revenue = float(b["rate"] or 0) * nights
        estimated_revenue += revenue
        overlapping_bookings += 1
        rs = room_stats[b["room_id"]]
        rs["sold_nights"] += nights
        rs["revenue"] += revenue
        rs["bookings"] += 1
        if not rs["out_of_order"]:
            sold_nights += nights
        bs = building_stats.setdefault(b["building"], {"building": b["building"], "sold_nights": 0, "revenue": 0.0, "bookings": 0})
        if not rs["out_of_order"]:
            bs["sold_nights"] += nights
        bs["revenue"] += revenue
        bs["bookings"] += 1
    capacity_nights = sum(1 for r in rooms if not r["out_of_order"]) * period_days
    rooms_sold = sum(1 for rs in room_stats.values() if rs["sold_nights"] > 0 and not rs["out_of_order"])
    occupancy = (sold_nights / capacity_nights * 100) if capacity_nights else 0
    for rs in room_stats.values():
        rs["occupancy"] = (rs["sold_nights"] / period_days * 100) if period_days and not rs["out_of_order"] else 0
    for bs in building_stats.values():
        building_room_count = sum(1 for r in rooms if r["building"] == bs["building"] and not r["out_of_order"])
        building_capacity = building_room_count * period_days
        bs["occupancy"] = (bs["sold_nights"] / building_capacity * 100) if building_capacity else 0
        bs["room_count"] = building_room_count
    audit("view", "report", None, {"type": "occupancy", "start": start_date.isoformat(), "end": end_date.isoformat()})
    return render_template("reports.html", start_date=start_date, end_date=end_date, period_days=period_days,
                           rooms_count=len(rooms), capacity_nights=capacity_nights, sold_nights=sold_nights,
                           rooms_sold=rooms_sold, occupancy=occupancy, estimated_revenue=estimated_revenue,
                           overlapping_bookings=overlapping_bookings,
                           room_stats=sorted(room_stats.values(), key=lambda x: (x["building"], x["room_name"])),
                           building_stats=sorted(building_stats.values(), key=lambda x: x["building"]))

def _report_export_data(start_date, end_date):
    """Build normalized report tables suitable for formulas, filters and pivot charts."""
    end_exclusive = end_date + timedelta(days=1)
    period_days = (end_exclusive - start_date).days
    conn = connect()
    rooms = conn.execute("""SELECT id, name, building, out_of_order FROM rooms WHERE active=1
        ORDER BY CASE building WHEN 'Corp A' THEN 1 WHEN 'Corp B' THEN 2 WHEN 'Corp C' THEN 3 ELSE 4 END, sort_order, name""").fetchall()
    bookings = conn.execute("""
        SELECT b.*, r.name AS room_name, r.building AS building, r.out_of_order AS room_defective
        FROM bookings b JOIN rooms r ON r.id=b.room_id
        WHERE r.active=1 AND b.status IN ('confirmed','checked_in','checked_out')
          AND b.check_in < ? AND b.check_out > ?
        ORDER BY b.check_in, r.building, r.sort_order, r.name, b.id
    """, (end_exclusive.isoformat(), start_date.isoformat())).fetchall()
    all_bookings = [dict(b) for b in bookings]
    conn.close()
    room_stats = {r['id']: {'room_id':r['id'], 'room_name':r['name'], 'building':r['building'],
        'state':'Defectă' if r['out_of_order'] else 'Utilizabilă', 'out_of_order':bool(r['out_of_order']),
        'sold_nights':0, 'revenue':0.0, 'bookings':0} for r in rooms}
    building_stats = {}
    daily = { (start_date + timedelta(days=i)).isoformat(): {'date':(start_date+timedelta(days=i)).isoformat(),
        'sold_nights':0, 'occupied_rooms':0, 'arrivals':0, 'departures':0, 'revenue':0.0} for i in range(period_days)}
    reservation_rows = []
    for b in all_bookings:
        ci=date.fromisoformat(b['check_in']); co=date.fromisoformat(b['check_out'])
        clipped_start=max(ci,start_date); clipped_end=min(co,end_exclusive)
        nights=max(0,(clipped_end-clipped_start).days)
        if not nights: continue
        rate=float(b['rate'] or 0); revenue=rate*nights
        rs=room_stats[b['room_id']]
        rs['sold_nights']+=nights; rs['revenue']+=revenue; rs['bookings']+=1
        bs=building_stats.setdefault(b['building'], {'building':b['building'],'sold_nights':0,'revenue':0.0,'bookings':0})
        if not rs['out_of_order']:
            bs['sold_nights']+=nights
        bs['revenue']+=revenue; bs['bookings']+=1
        for day in range((clipped_end-clipped_start).days):
            d=(clipped_start+timedelta(days=day)).isoformat()
            daily[d]['revenue']+=rate
            if not rs['out_of_order']:
                daily[d]['sold_nights']+=1
                daily[d]['occupied_rooms']+=1
        if start_date <= ci < end_exclusive: daily[ci.isoformat()]['arrivals']+=1
        if start_date <= co < end_exclusive: daily[co.isoformat()]['departures']+=1
        reservation_rows.append([b['id'], b['guest_name'], b['phone'] or '', b['room_name'], b['building'],
            b['check_in'], b['check_out'], nights, rate, round(revenue,2), b['status'], b['source'] or '',
            b['receipt_number'] or '', b['group_id'] or ''])
    usable_rooms=sum(1 for r in rooms if not r['out_of_order'])
    capacity=usable_rooms*period_days
    sold=sum(x['sold_nights'] for x in room_stats.values() if not x['out_of_order'])
    revenue_total=sum(x['revenue'] for x in room_stats.values())
    summary = [
        ['Indicator','Valoare'],
        ['Perioada',f'{start_date.isoformat()} – {end_date.isoformat()}'],
        ['Camere în inventar',len(rooms)],
        ['Camere utilizabile',usable_rooms],
        ['Zile incluse',period_days],
        ['Nopți disponibile',capacity],
        ['Nopți vândute',sold],
        ['Grad de ocupare (%)', round(sold/capacity*100,2) if capacity else 0],
        ['Camere utilizabile cu nopți vândute',sum(1 for x in room_stats.values() if x['sold_nights']>0 and not x['out_of_order'])],
        ['Rezervări confirmate în perioadă',len(reservation_rows)],
        ['Venit estimat (lei)',round(revenue_total,2)]
    ]
    buildings=[['Corp','Camere utilizabile','Nopți vândute','Capacitate nopți','Grad ocupare (%)','Rezervări','Venit estimat (lei)']]
    for name in ['Corp A','Corp B','Corp C'] + sorted({r['building'] for r in rooms}-{ 'Corp A','Corp B','Corp C'}):
        room_count=sum(1 for r in rooms if r['building']==name and not r['out_of_order'])
        bs=building_stats.get(name,{'sold_nights':0,'revenue':0,'bookings':0})
        buildings.append([name,room_count,bs['sold_nights'],room_count*period_days,
            round(bs['sold_nights']/(room_count*period_days)*100,2) if room_count and period_days else 0,
            bs['bookings'],round(bs['revenue'],2)])
    details=[['Corp','Cameră','Stare','Nopți vândute','Nopți disponibile','Grad ocupare (%)','Rezervări','Venit estimat (lei)']]
    for r in rooms:
        rs=room_stats[r['id']]
        avail=0 if r['out_of_order'] else period_days
        details.append([r['building'],r['name'],rs['state'],rs['sold_nights'],avail,
            round(rs['sold_nights']/avail*100,2) if avail and not r['out_of_order'] else 0,rs['bookings'],round(rs['revenue'],2)])
    daily_rows=[['Data','Camere ocupate','Camere utilizabile','Sosiri','Plecări','Nopți vândute','Venit estimat (lei)','Grad ocupare (%)']]
    for d,vals in daily.items():
        daily_rows.append([d,vals['occupied_rooms'],usable_rooms,vals['arrivals'],vals['departures'],
            vals['sold_nights'],round(vals['revenue'],2),
            round(vals['occupied_rooms']/usable_rooms*100,2) if usable_rooms else 0])
    reservation_headers=['ID rezervare','Client / grup','Telefon','Cameră','Corp','Check-in','Check-out','Nopți în perioadă','Tarif/noapte (lei)','Venit estimat (lei)','Stare','Sursă','Chitanță','ID grup']
    reservations=[reservation_headers]+reservation_rows
    return {'summary':summary,'buildings':buildings,'details':details,'daily':daily_rows,'reservations':reservations}

def _report_csv_response(start_date, end_date, fallback=False):
    import csv
    from io import StringIO
    from flask import Response
    data = _report_export_data(start_date, end_date)
    stream = StringIO(newline='')
    writer = csv.writer(stream, delimiter=';')
    # CSV stays rectangular/normalized; separate CSV exports can be generated from XLS sheets.
    writer.writerows([[safety.csv_safe(c) for c in row] for row in data['reservations']])
    payload = ('\ufeff' + stream.getvalue()).encode('utf-8')
    filename = f"raport_popas_caprioara_rezervari_{start_date.isoformat()}_{end_date.isoformat()}.csv"
    audit('export_fallback', 'report', None, {'format':'csv-fallback' if fallback else 'csv',
        'start':start_date.isoformat(),'end':end_date.isoformat()})
    return Response(payload, mimetype='text/csv; charset=utf-8',
        headers={'Content-Disposition':f'attachment; filename={filename}',
                 'X-Report-Fallback':'csv' if fallback else 'csv'})


@app.route("/admin/reports/export.csv")
@admin_required
def reports_export_csv():
    start_raw = request.args.get("start", (clock.today() - timedelta(days=29)).isoformat())
    end_raw = request.args.get("end", clock.today().isoformat())
    try:
        start_date = date.fromisoformat(start_raw)
        end_date = date.fromisoformat(end_raw)
    except ValueError:
        abort(400, description="Perioada raportului conține o dată invalidă.")
    if end_date < start_date:
        abort(400, description="Data finală trebuie să fie egală sau după data de început.")
    return _report_csv_response(start_date, end_date)


@app.route("/admin/reports/export")
@admin_required
def reports_export():
    """Export a formatted XLSX workbook with formulas, filters, and occupancy charts."""
    start_raw = request.args.get("start", (clock.today() - timedelta(days=29)).isoformat())
    end_raw = request.args.get("end", clock.today().isoformat())
    try:
        start_date = date.fromisoformat(start_raw)
        end_date = date.fromisoformat(end_raw)
    except ValueError:
        abort(400, description="Perioada raportului conține o dată invalidă.")
    if end_date < start_date:
        abort(400, description="Data finală trebuie să fie egală sau după data de început.")
    try:
        from io import BytesIO
        from flask import send_file
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.chart.label import DataLabelList
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
        data = _report_export_data(start_date, end_date)
        wb = Workbook()
        wb.remove(wb.active)
        navy = "173D35"; green = "26745C"; pale = "EAF4EF"; light = "F5F8F6"; white = "FFFFFF"
        thin = Side(style="thin", color="D9E3DD")
        numeric_columns = {
            'Rezumat': {2,3,4,5,6,7,8,9,10,11},
            'Pe corpuri': {2,3,4,5,6,7},
            'Pe camere': {4,5,6,7,8},
            'Date zilnice': {2,3,4,5,6,7,8},
            'Rezervari': {1,8,9,10},
        }
        sheet_map = [('Rezumat','summary'),('Pe corpuri','buildings'),('Pe camere','details'),('Date zilnice','daily'),('Rezervari','reservations')]
        for sheet_name, key in sheet_map:
            rows = data[key]
            ws = wb.create_sheet(sheet_name)
            for row in rows:
                ws.append([safety.csv_safe(c) for c in row])
            ws.freeze_panes = 'A2'
            ws.sheet_view.showGridLines = False
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color=white, size=10)
                cell.fill = PatternFill('solid', fgColor=navy)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                cell.border = Border(bottom=Side(style='medium', color=green))
            ws.row_dimensions[1].height = 30
            for ri in range(2, ws.max_row + 1):
                if ri % 2 == 0:
                    for cell in ws[ri]: cell.fill = PatternFill('solid', fgColor=light)
                for ci in range(1, ws.max_column + 1):
                    cell = ws.cell(ri, ci)
                    cell.alignment = Alignment(vertical='top')
                    cell.border = Border(bottom=thin)
                    if isinstance(cell.value, (int, float)):
                        cell.alignment = Alignment(horizontal='right', vertical='top')
                        if sheet_name == 'Rezervari' and ci in (9,10): cell.number_format = '#,##0.00'
                        elif sheet_name == 'Pe corpuri' and ci == 7: cell.number_format = '#,##0.00'
                        elif sheet_name == 'Pe camere' and ci == 8: cell.number_format = '#,##0.00'
                        elif sheet_name == 'Date zilnice' and ci == 7: cell.number_format = '#,##0.00'
                        elif ((sheet_name == 'Pe corpuri' and ci == 5) or (sheet_name == 'Pe camere' and ci == 6) or (sheet_name == 'Date zilnice' and ci == 8)):
                            cell.number_format = '0.0"%"'
                        else: cell.number_format = '#,##0.0' if isinstance(cell.value,float) else '#,##0'
            for ci in range(1, ws.max_column + 1):
                values = [str(ws.cell(ri,ci).value or '') for ri in range(1, min(ws.max_row, 250) + 1)]
                max_len = max((len(v) for v in values), default=10)
                ws.column_dimensions[get_column_letter(ci)].width = min(max(max_len + 2, 12), 36)
            if ws.max_row > 1 and ws.max_column > 1:
                safe_name = 'Tbl' + ''.join(ch for ch in sheet_name if ch.isalnum())[:20]
                table = Table(displayName=safe_name, ref=ws.dimensions)
                table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showRowStripes=True, showColumnStripes=False)
                ws.add_table(table)
        # Make summary calculations auditable and formula-driven.
        ws = wb['Rezumat']
        ws['B8'] = '=IF(B6=0,0,B7/B6*100)'
        ws['B8'].number_format = '0.0"%"'
        ws['B8'].font = Font(bold=True, color=green)
        ws['A13'] = 'Notă metodologică'
        ws['B13'] = 'Grad ocupare = nopți vândute / nopți disponibile × 100. Camerele marcate defecte sunt excluse din capacitate; verificați rezervările vechi pe camere defecte.'
        ws['A13'].font = Font(bold=True, color=navy)
        ws['B13'].alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[13].height = 42
        ws.column_dimensions['A'].width = 34
        ws.column_dimensions['B'].width = 78
        # Formula-based occupancy in summary and grouped reports, so users can adjust values later.
        ws_build = wb['Pe corpuri']
        for row in range(2, ws_build.max_row + 1):
            ws_build.cell(row, 5, f'=IF(D{row}=0,0,C{row}/D{row}*100)')
            ws_build.cell(row, 5).number_format = '0.0"%"'
        ws_room = wb['Pe camere']
        for row in range(2, ws_room.max_row + 1):
            ws_room.cell(row, 6, f'=IF(E{row}=0,0,D{row}/E{row}*100)')
            ws_room.cell(row, 6).number_format = '0.0"%"'
        ws_daily = wb['Date zilnice']
        for row in range(2, ws_daily.max_row + 1):
            ws_daily.cell(row, 8, f'=IF(C{row}=0,0,B{row}/C{row}*100)')
            ws_daily.cell(row, 8).number_format = '0.0"%"'
        # Occupancy charts are placed on the summary sheet, backed by the detailed tables.
        bar = BarChart()
        bar.type = 'bar'; bar.style = 10; bar.title = 'Grad de ocupare pe corp (%)'; bar.y_axis.title = 'Corp'; bar.x_axis.title = 'Grad ocupare (%)'
        bar.height = 7; bar.width = 13
        bar.add_data(Reference(ws_build, min_col=5, min_row=1, max_row=ws_build.max_row), titles_from_data=True)
        bar.set_categories(Reference(ws_build, min_col=1, min_row=2, max_row=ws_build.max_row))
        bar.legend = None
        ws.add_chart(bar, 'D2')
        line = LineChart(); line.style = 13; line.title = 'Grad de ocupare zilnic (%)'; line.y_axis.title = 'Grad ocupare (%)'; line.x_axis.title = 'Data'; line.height = 7; line.width = 18
        line.add_data(Reference(ws_daily, min_col=8, min_row=1, max_row=ws_daily.max_row), titles_from_data=True)
        line.set_categories(Reference(ws_daily, min_col=1, min_row=2, max_row=ws_daily.max_row))
        line.legend = None
        ws.add_chart(line, 'D17')
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = 'auto'
        output = BytesIO(); wb.save(output); output.seek(0)
        audit('export', 'report', None, {'format':'xlsx', 'start':start_date.isoformat(), 'end':end_date.isoformat()})
        filename = f"raport_popas_caprioara_{start_date.isoformat()}_{end_date.isoformat()}.xlsx"
        return send_file(output, as_attachment=True, download_name=filename,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as exc:
        app.logger.exception("Exportul XLSX a eșuat; se livrează CSV de rezervă.")
        audit('export_fallback', 'report', None, {'requested_format':'xlsx','error':str(exc)[:300],
            'start':start_date.isoformat(),'end':end_date.isoformat()})
        return _report_csv_response(start_date, end_date, fallback=True)

@app.route("/booking/<int:booking_id>/oblio-test", methods=["POST"])
@login_required
def oblio_test_invoice(booking_id):
    conn = connect()
    booking = conn.execute("SELECT b.*, r.name AS room_name FROM bookings b JOIN rooms r ON r.id=b.room_id WHERE b.id=?", (booking_id,)).fetchone()
    if not booking:
        conn.close()
        abort(404)
    conn.execute("BEGIN IMMEDIATE")
    existing = conn.execute("SELECT * FROM invoices WHERE booking_id=?", (booking_id,)).fetchone()
    key = hashlib.sha256(f"popas-caprioara-booking-{booking_id}-invoice-v1".encode()).hexdigest()
    if existing:
        conn.execute("INSERT INTO invoice_attempts(invoice_id, booking_id, idempotency_key, status, message, attempted_by) VALUES (?, ?, ?, 'duplicate_blocked', ?, ?)", (existing["id"], booking_id, key, "Cererea repetată a fost blocată; nu s-a creat o a doua factură de test.", g.user["id"]))
        conn.commit(); conn.close()
        audit("invoice_duplicate_blocked", "booking", booking_id, {"invoice_id": existing["id"]})
        flash("Protecție anti-duplicat: există deja o înregistrare de factură pentru această rezervare. Nu a fost creată alta.", "error")
        return redirect(url_for("edit_booking", booking_id=booking_id))
    total = round(float(booking["rate"] or 0) * max(0, (date.fromisoformat(booking["check_out"]) - date.fromisoformat(booking["check_in"])).days), 2)
    # Local test mode only: no request is sent to Oblio and no fiscal document is issued.
    cur = conn.execute("INSERT INTO invoices(booking_id, provider, mode, status, idempotency_key, issue_date, total, response_json, created_by) VALUES (?, 'oblio', 'test', 'test_emitted', ?, ?, ?, ?, ?)", (booking_id, key, clock.today().isoformat(), total, json.dumps({"mode":"local_test","real_api_call":False,"guest_name":booking["guest_name"],"room":booking["room_name"]}, ensure_ascii=False), g.user["id"]))
    invoice_id = cur.lastrowid
    conn.execute("INSERT INTO invoice_attempts(invoice_id, booking_id, idempotency_key, status, message, response_json, attempted_by) VALUES (?, ?, ?, 'test_emitted', ?, ?, ?)", (invoice_id, booking_id, key, "Simulare locală: nu s-a apelat API-ul Oblio și nu s-a emis document fiscal.", json.dumps({"total":total}), g.user["id"]))
    conn.commit(); conn.close()
    audit("invoice_test_emitted", "invoice", invoice_id, {"booking_id": booking_id, "total": total, "real_api_call": False})
    flash(f"Test Oblio reușit: înregistrare locală #{invoice_id}, total calculat {total:.2f} lei. Nu s-a emis o factură fiscală reală.", "success")
    return redirect(url_for("edit_booking", booking_id=booking_id))

@app.template_filter("money")
def money(value):
    return f"{float(value or 0):.2f}".replace(".", ",")

@app.route("/")
@login_required
def index():
    today = request.args.get("start", clock.today().isoformat())
    try:
        start_date = date.fromisoformat(today)
    except ValueError:
        start_date = clock.today()
    days = max(7, min(21, int(request.args.get("days", 14))))
    dates = [start_date + timedelta(days=i) for i in range(days)]
    end_date = dates[-1] + timedelta(days=1)
    conn = connect()
    rooms = conn.execute("SELECT * FROM rooms WHERE active=1 ORDER BY sort_order, name").fetchall()
    bookings = conn.execute("""
        SELECT b.*, r.name AS room_name FROM bookings b
        JOIN rooms r ON r.id=b.room_id
        WHERE b.status != 'cancelled' AND b.check_in < ? AND b.check_out > ?
        ORDER BY b.check_in, r.name
    """, (end_date.isoformat(), start_date.isoformat())).fetchall()
    calendar_users = conn.execute("SELECT id, display_name FROM users WHERE active=1 ORDER BY id").fetchall()
    search_term = request.args.get("q", "").strip()
    upcoming_sql = """
        SELECT b.*, r.name AS room_name FROM bookings b
        JOIN rooms r ON r.id=b.room_id
        WHERE b.status != 'cancelled' AND b.check_out >= ?
    """
    upcoming_params = [clock.today().isoformat()]
    if search_term:
        upcoming_sql += " AND (b.guest_name LIKE ? OR b.phone LIKE ? OR r.name LIKE ? OR b.email LIKE ?)"
        wildcard = f"%{search_term}%"
        upcoming_params.extend([wildcard, wildcard, wildcard, wildcard])
    upcoming_sql += " ORDER BY b.check_in, r.name LIMIT 100"
    all_upcoming = conn.execute(upcoming_sql, upcoming_params).fetchall()

    # Daily reception dashboard: every list is derived from the live database.
    today_iso = clock.today().isoformat()
    housekeeping_rooms = conn.execute("SELECT id, name, building FROM rooms WHERE active=1 AND out_of_order=0 ORDER BY CASE building WHEN 'Corp A' THEN 1 WHEN 'Corp B' THEN 2 WHEN 'Corp C' THEN 3 ELSE 4 END, sort_order, name").fetchall()
    housekeeping_tasks = conn.execute("""SELECT h.*, r.name AS room_name, r.building FROM housekeeping_requests h JOIN rooms r ON r.id=h.room_id
        WHERE h.service_date=? AND h.status IN ('open','done') ORDER BY CASE h.status WHEN 'open' THEN 0 ELSE 1 END, r.building, r.sort_order, r.name, h.id DESC""", (today_iso,)).fetchall()
    housekeeping_task_items = []
    for row in housekeeping_tasks:
        item = dict(row)
        try:
            item['requested_time'] = datetime.strptime(item.get('created_at', ''), '%Y-%m-%d %H:%M:%S').replace(tzinfo=__import__('datetime').timezone.utc).astimezone().strftime('%H:%M')
        except (ValueError, TypeError):
            item['requested_time'] = (item.get('created_at') or '')[-8:-3]
        housekeeping_task_items.append(item)
    room_clean_status = {r['room_id']: r for r in conn.execute("SELECT room_id, is_clean, updated_at, updated_by, updated_by_name FROM room_cleaning WHERE service_date=?", (today_iso,)).fetchall()}
    arrivals_today = conn.execute("""
        SELECT b.*, r.name AS room_name FROM bookings b JOIN rooms r ON r.id=b.room_id
        WHERE b.check_in=? AND b.status IN ('confirmed','provisional','unconfirmed')
        ORDER BY CASE b.status WHEN 'unconfirmed' THEN 0 WHEN 'provisional' THEN 1 ELSE 2 END, r.name
    """, (today_iso,)).fetchall()
    departures_today = conn.execute("""
        SELECT b.*, r.name AS room_name FROM bookings b JOIN rooms r ON r.id=b.room_id
        WHERE b.check_out=? AND b.status IN ('confirmed','checked_in','checked_out')
        ORDER BY r.name
    """, (today_iso,)).fetchall()
    in_house_today = conn.execute("""
        SELECT b.*, r.name AS room_name FROM bookings b JOIN rooms r ON r.id=b.room_id
        WHERE b.check_in<=? AND b.check_out>? AND b.status IN ('confirmed','checked_in')
        ORDER BY r.name
    """, (today_iso, today_iso)).fetchall()
    reconfirmation_today = conn.execute("""
        SELECT b.*, r.name AS room_name FROM bookings b JOIN rooms r ON r.id=b.room_id
        WHERE b.check_in=? AND b.source='Direct' AND b.group_id IS NULL
          AND b.status IN ('confirmed','provisional','unconfirmed')
          AND b.reconfirmed_at IS NULL
          AND date(b.created_at, 'localtime')=?
        ORDER BY CASE b.status WHEN 'unconfirmed' THEN 0 ELSE 1 END, r.name
    """, (today_iso, today_iso)).fetchall()
    occupied_room_ids = conn.execute("""
        SELECT DISTINCT b.room_id FROM bookings b JOIN rooms r ON r.id=b.room_id AND r.active=1
        WHERE b.check_in<=? AND b.check_out>? AND b.status NOT IN ('cancelled','unconfirmed','checked_out')
    """, (today_iso, today_iso)).fetchall()
    conn.close()
    by_room = {r["id"]: [] for r in rooms}
    for b in bookings:
        by_room[b["room_id"]].append(dict(b))
    total_rooms = len(rooms)
    defective_rooms = sum(1 for r in rooms if r["out_of_order"])
    occupied_today = len({b["room_id"] for b in in_house_today})
    arrivals = len(arrivals_today)
    departures = len(departures_today)
    occupied_ids = {row[0] for row in occupied_room_ids}
    available_today = sum(1 for room in rooms if not room["out_of_order"] and room["id"] not in occupied_ids)

    def dashboard_row(row):
        item = dict(row)
        nights = max(0, (date.fromisoformat(item["check_out"]) - date.fromisoformat(item["check_in"])).days)
        item["estimated_balance"] = max(0.0, float(item.get("rate") or 0) * nights - float(item.get("deposit") or 0))
        item["nights"] = nights
        return item

    arrivals_today = [dashboard_row(row) for row in arrivals_today]
    departures_today = [dashboard_row(row) for row in departures_today]
    in_house_today = [dashboard_row(row) for row in in_house_today]
    occupied_today_ids = {b["room_id"] for b in in_house_today}
    reconfirmation_today = [dashboard_row(row) for row in reconfirmation_today]
    legend_by_color = {}
    for u in calendar_users:
        legend_by_color[color_for_user(u["id"])] = {"name": u["display_name"], "color": color_for_user(u["id"])}
    for b in bookings:
        color = b["created_by_color"] or "#71817c"
        legend_by_color.setdefault(color, {"name": b["created_by_name"] or "Importat / necunoscut", "color": color})
    calendar_legend = list(legend_by_color.values())

    # Calendar rows are explicitly numbered so each building can have a divider
    # while booking bars remain aligned with the physical room row.
    calendar_rows = []
    grid_row = 2  # row 1 is the calendar header
    for building in ("Corp A", "Corp B", "Corp C"):
        building_rooms = [r for r in rooms if r["building"] == building]
        if not building_rooms:
            continue
        calendar_rows.append({"kind": "building", "building": building, "row_index": grid_row,
                              "room_count": len(building_rooms)})
        grid_row += 1
        for room in building_rooms:
            calendar_rows.append({"kind": "room", "room": room, "row_index": grid_row})
            grid_row += 1

    board_conn = connect()
    room_states = room_state.compute_room_states(board_conn)
    board_conn.close()
    is_admin = g.user["role"] == "admin"
    for item in room_states:
        if item["booking_id"]:
            item["href"] = url_for("edit_booking", booking_id=item["booking_id"])
        elif item["state"] == "out_of_order":
            item["href"] = url_for("room_maintenance") if is_admin else None
        else:
            item["href"] = url_for("new_booking", room_id=item["room_id"], check_in=clock.today().isoformat())
    room_state_counts = room_state.state_counts(room_states)
    room_state_map = {item["room_id"]: item for item in room_states}
    board_conn = connect()
    recent_bookings = board_conn.execute("""SELECT b.*, r.name AS room_name FROM bookings b JOIN rooms r ON r.id=b.room_id
        ORDER BY b.updated_at DESC, b.id DESC LIMIT 3""").fetchall()
    dc_info = deep_clean.due_info(board_conn, room_states)
    night_pending = board_conn.execute("""SELECT b.id, b.guest_name, b.rate, r.name AS room_name FROM bookings b JOIN rooms r ON r.id=b.room_id
        WHERE b.night_payment='morning' AND b.status IN ('checked_in','confirmed') AND b.check_in >= ? ORDER BY b.night_at""",
        ((clock.today() - timedelta(days=2)).isoformat(),)).fetchall()
    board_conn.close()

    return render_template("index.html", room_states=room_states, room_state_map=room_state_map, recent_bookings=recent_bookings, dc=dc_info, night_pending=night_pending, room_state_counts=room_state_counts,
                           room_state_order=room_state.STATE_ORDER, room_state_labels=room_state.STATE_LABELS, rooms=rooms, dates=dates, by_room=by_room,
                           calendar_rows=calendar_rows, start=start_date, days=days, date=date,
                           timedelta=timedelta, today=clock.today().isoformat(), bookings=all_upcoming,
                           total_rooms=total_rooms, defective_rooms=defective_rooms, occupied_today=occupied_today,
                           arrivals=arrivals, departures=departures, calendar_legend=calendar_legend,
                           arrivals_today=arrivals_today, departures_today=departures_today,
                           in_house_today=in_house_today, reconfirmation_today=reconfirmation_today,
                           available_today=available_today, search_term=search_term,
                           housekeeping_rooms=housekeeping_rooms, housekeeping_tasks=housekeeping_task_items,
                           room_clean_status=room_clean_status, occupied_today_ids=occupied_today_ids)

@app.route("/booking/new", methods=["GET", "POST"])
@login_required
def new_booking():
    conn = connect()
    rooms = conn.execute("SELECT * FROM rooms WHERE active=1 AND out_of_order=0 ORDER BY sort_order, name").fetchall()
    if request.method == "POST":
        try:
            data = get_form_data(request.form)
        except (TypeError, ValueError):
            data = get_form_data_safe(request.form)
            flash("Verificați valorile numerice introduse.", "error")
            conn.close()
            return render_template("booking_form.html", rooms=rooms, data=data, mode="new", calendar_start=data.get("check_in") or clock.today().isoformat())
        error = validate(data, {r["id"] for r in rooms})
        if error:
            flash(error, "error")
            conn.close()
            return render_template("booking_form.html", rooms=rooms, data=data, mode="new", calendar_start=data.get("check_in") or clock.today().isoformat())
        try:
            conn.execute("BEGIN IMMEDIATE")
            conflict = overlaps(conn, data["room_id"], data["check_in"], data["check_out"])
            if conflict:
                conn.rollback()
                flash(f"Conflict: camera este deja rezervată pentru {conflict['guest_name']} "
                      f"({conflict['check_in']} – {conflict['check_out']}). Alegeți altă cameră sau perioadă.", "error")
                conn.close()
                return render_template("booking_form.html", rooms=rooms, data=data, mode="new", calendar_start=data.get("check_in") or clock.today().isoformat())
            conn.execute("""
                INSERT INTO bookings
                (guest_name, phone, email, room_id, check_in, check_out, guests, rate, deposit, source, status, notes, late_arrival, night_note,
                 created_by, created_by_name, created_by_color)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, tuple(data[k] for k in ["guest_name","phone","email","room_id","check_in","check_out","guests","rate","deposit","source","status","notes","late_arrival","night_note"])
                 + (g.user["id"], g.user["display_name"], color_for_user(g.user["id"])))
            booking_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
            conn.commit()
            conn.close()
            audit("create", "booking", booking_id, {"room_id": data["room_id"], "check_in": data["check_in"], "check_out": data["check_out"], "status": data["status"]})
            flash("Rezervarea a fost salvată. Disponibilitatea a fost verificată pe server.", "success")
            return redirect(url_for("index", start=data["check_in"]))
        except Exception:
            conn.rollback()
            conn.close()
            raise
    conn.close()
    today = clock.today().isoformat()
    requested_date = request.args.get("check_in", request.args.get("date", today))
    try:
        selected_date = date.fromisoformat(requested_date)
    except ValueError:
        selected_date = clock.today()
    requested_room = request.args.get("room_id", "")
    if requested_room and not any(str(r["id"]) == str(requested_room) for r in rooms):
        requested_room = ""
    return render_template("booking_form.html", rooms=rooms, data={
        "guest_name":"","phone":"","email":"","room_id":requested_room,"check_in":selected_date.isoformat(),
        "check_out":(selected_date+timedelta(days=1)).isoformat(),"guests":"","late_arrival":0,"night_note":"",
        "rate":"","deposit":0,"source":"Direct","status":"confirmed","notes":"","receipt_number":"","group_id":None
    }, mode="new", calendar_start=selected_date.isoformat())

@app.route("/booking/group/new", methods=["GET", "POST"])
@login_required
def new_group_booking():
    conn = connect()
    # Initially show active rooms; the date-aware API filters out occupied rooms.
    rooms = conn.execute("SELECT * FROM rooms WHERE active=1 AND out_of_order=0 ORDER BY sort_order, name").fetchall()
    if request.method == "POST":
        guest_name = request.form.get("guest_name", "").strip()
        phone = request.form.get("phone", "").strip()
        email = request.form.get("email", "").strip()
        check_in = request.form.get("check_in", "")
        check_out = request.form.get("check_out", "")
        source = request.form.get("source", "Direct")
        notes = request.form.get("notes", "").strip()
        try:
            room_ids = list(dict.fromkeys(int(x) for x in request.form.getlist("room_ids")))
            guests = int((request.form.get("guests") or "").strip() or 0)
            rate = float(request.form.get("rate", "0"))
            deposit = float(request.form.get("deposit", "0"))
            start, end = date.fromisoformat(check_in), date.fromisoformat(check_out)
        except (ValueError, TypeError):
            room_ids, guests, rate, deposit, start, end = [], 0, 0, 0, None, None
        valid_ids = {r['id'] for r in rooms}
        error = None
        if not guest_name: error = "Introduceți numele grupului sau al persoanei de contact."
        elif not phone: error = "Numărul de telefon al persoanei de contact este obligatoriu pentru rezervările de grup."
        elif start is None or end is None or end <= start: error = "Verificați datele de sosire și plecare."
        elif not room_ids or any(rid not in valid_ids for rid in room_ids): error = "Selectați cel puțin o cameră disponibilă pentru perioada aleasă."
        elif guests < 1: error = "Numărul de persoane per cameră este obligatoriu: alegeți 1 sau 2."
        elif guests > 2: error = "Numărul de persoane per cameră trebuie să fie între 1 și 2."
        elif not (0 < rate <= 10000): error = "Tariful pe noapte este obligatoriu: mai mare decât 0 și cel mult 10.000 lei."
        elif not (0 <= deposit <= 100000): error = "Avansul per cameră trebuie să fie între 0 și 100.000 lei."
        elif len(guest_name) > 120 or len(notes) > 1000: error = "Numele (maximum 120) sau observațiile (maximum 1000) sunt prea lungi."
        elif not safety.PHONE_RE.fullmatch(phone): error = "Telefon invalid: folosiți cifre, spații și caracterele + ( ) - . /"
        elif email and not safety.EMAIL_RE.fullmatch(email): error = "Adresa de e-mail nu este validă."
        elif source not in SOURCES: error = "Sursa rezervării nu este validă."
        elif (end - start).days > 90: error = "Șederea nu poate depăși 90 de nopți."
        if not error:
            conn.execute("BEGIN IMMEDIATE")
            conflicts = []
            for room_id in room_ids:
                room_state = conn.execute("SELECT name, out_of_order FROM rooms WHERE id=? AND active=1", (room_id,)).fetchone()
                if not room_state or room_state['out_of_order']:
                    conflicts.append(f"camera {room_id} nu mai este disponibilă sau este marcată ca defectă")
                    continue
                conflict = overlaps(conn, room_id, check_in, check_out)
                if conflict:
                    room_name = room_state['name']
                    conflicts.append(f"camera {room_name} ({conflict['guest_name']}, {conflict['check_in']}–{conflict['check_out']})")
            if conflicts:
                conn.rollback()
                error = "Nu s-a creat grupul. Camere indisponibile: " + "; ".join(conflicts)
            else:
                group_id = secrets.token_hex(8)
                for room_id in room_ids:
                    conn.execute("""INSERT INTO bookings
                        (guest_name, phone, email, room_id, check_in, check_out, guests, rate, deposit, source, status, notes,
                         created_by, created_by_name, created_by_color, group_id, receipt_number)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'provisional', ?, ?, ?, ?, ?, '')""",
                        (guest_name, phone, email, room_id, check_in, check_out, guests, rate, deposit, source, notes,
                         g.user['id'], g.user['display_name'], color_for_user(g.user['id']), group_id))
                conn.commit()
                booking_ids = [r['id'] for r in conn.execute("SELECT id FROM bookings WHERE group_id=?", (group_id,)).fetchall()]
                conn.close()
                audit("create_group", "booking_group", group_id, {"booking_ids": booking_ids, "rooms": room_ids, "check_in": check_in, "check_out": check_out, "status": "provisional"})
                flash(f"Rezervarea de grup a fost creată provizoriu pentru {len(room_ids)} camere. Pentru confirmare este obligatoriu numărul chitanței.", "success")
                return redirect(url_for("index", start=check_in))
        flash(error, "error")
    conn.close()
    today = clock.today().isoformat()
    form_start_raw = request.form.get("check_in", request.args.get("check_in", today))
    try:
        form_start = date.fromisoformat(form_start_raw)
    except (ValueError, TypeError):
        form_start = clock.today()
        form_start_raw = form_start.isoformat()
    form_end_default = (form_start + timedelta(days=1)).isoformat()
    return render_template("group_booking_form.html", rooms=rooms, today=today,
                           data={"guest_name":request.form.get("guest_name", ""), "phone":request.form.get("phone", ""),
                                 "email":request.form.get("email", ""), "check_in":form_start_raw,
                                 "check_out":request.form.get("check_out", request.args.get("check_out", form_end_default)),
                                 "room_ids":request.form.getlist("room_ids"), "guests":request.form.get("guests", ""),
                                 "rate":request.form.get("rate", ""), "deposit":request.form.get("deposit", "0"),
                                 "source":request.form.get("source", "Direct"), "notes":request.form.get("notes", "")})


def get_form_data(form):
    return {
        "guest_name": form.get("guest_name", "").strip(),
        "phone": form.get("phone", "").strip(),
        "email": form.get("email", "").strip(),
        "room_id": int(form.get("room_id") or 0),
        "check_in": form.get("check_in", ""),
        "check_out": form.get("check_out", ""),
        "guests": int((form.get("guests") or "").strip() or 0),   # gol = 0 -> validarea cere explicit numărul de persoane
        "rate": float(form.get("rate") or 0),
        "deposit": float(form.get("deposit") or 0),
        "receipt_number": form.get("receipt_number", "").strip(),
        "source": form.get("source", "Direct"),
        "status": form.get("status", "confirmed"),
        "notes": form.get("notes", "").strip(),
        "late_arrival": 1 if form.get("late_arrival") == "1" else 0,
        "night_note": form.get("night_note", "").strip()[:300],
    }

def get_form_data_safe(form):
    """Return redisplay-safe form values if a numeric field is malformed."""
    data = {key: form.get(key, "") for key in [
        "guest_name", "phone", "email", "room_id", "check_in", "check_out",
        "guests", "rate", "deposit", "source", "status", "notes", "receipt_number", "night_note"]}
    data["late_arrival"] = 1 if form.get("late_arrival") == "1" else 0
    data["guest_name"] = (data["guest_name"] or "").strip()
    data["notes"] = (data["notes"] or "").strip()
    data["guests"] = data["guests"] or ""
    data["rate"] = data["rate"] or ""
    data["deposit"] = data["deposit"] or 0
    data["source"] = data["source"] or "Direct"
    data["status"] = data["status"] or "confirmed"
    return data

SOURCES = ["Direct", "Telefon", "WhatsApp", "Website", "Booking.com", "Expedia", "Altă platformă", "Agenție", "Recepție noapte"]


def validate(d, valid_room_ids=None):
    if not d["guest_name"]:
        return "Introduceți numele clientului."
    if len(d["guest_name"]) > 120:
        return "Numele clientului este prea lung (maximum 120 de caractere)."
    if d["phone"] and not safety.PHONE_RE.fullmatch(d["phone"]):
        return "Telefon invalid: folosiți cifre, spații și caracterele + ( ) - . /"
    if d["email"] and (len(d["email"]) > 120 or not safety.EMAIL_RE.fullmatch(d["email"])):
        return "Adresa de e-mail nu este validă."
    if len(d["notes"]) > 1000 or len(d.get("receipt_number", "")) > 40:
        return "Observațiile (maximum 1000 de caractere) sau numărul chitanței (maximum 40) sunt prea lungi."
    if d["source"] not in SOURCES:
        return "Sursa rezervării nu este validă."
    try:
        start = date.fromisoformat(d["check_in"])
        end = date.fromisoformat(d["check_out"])
    except (ValueError, TypeError):
        return "Datele de sosire și plecare sunt invalide."
    if end <= start:
        return "Data plecării trebuie să fie după data sosirii."
    if d["room_id"] < 1 or (valid_room_ids is not None and d["room_id"] not in valid_room_ids):
        return "Selectați o cameră validă."
    if d["guests"] < 1:
        return "Numărul de persoane în cameră este obligatoriu: alegeți 1 sau 2."
    if d["guests"] > 2:
        return "Capacitatea maximă este de 2 persoane pentru fiecare cameră."
    if not (0 < d["rate"] <= 10000):          # forma „not (a < x <= b)” respinge și NaN/infinit
        return "Tariful pe noapte este obligatoriu: mai mare decât 0 și cel mult 10.000 lei."
    if not (0 <= d["deposit"] <= 100000):
        return "Avansul trebuie să fie între 0 și 100.000 lei."
    if (end - start).days > 90:
        return "Șederea nu poate depăși 90 de nopți."
    if d["status"] not in STATUSES - {"cancelled"}:
        return "Statusul rezervării nu este valid."
    if d.get("night_note") and not night.night_note_allowed(d["status"], d.get("late_arrival"), d["check_in"], clock.today().isoformat()):
        return "Observațiile pentru recepția de noapte se pot adăuga doar pentru sosiri după ora 21:00. Bifați „Sosește după ora 21:00”."
    return None

@app.route("/booking/<int:booking_id>/edit", methods=["GET", "POST"])
@login_required
def edit_booking(booking_id):
    conn = connect()
    booking = conn.execute("SELECT * FROM bookings WHERE id=?", (booking_id,)).fetchone()
    if not booking:
        conn.close()
        flash("Rezervarea nu a fost găsită.", "error")
        return redirect(url_for("index"))
    group_bookings = []
    if booking["group_id"]:
        group_bookings = conn.execute("""SELECT b.id, b.status, r.name AS room_name FROM bookings b
            JOIN rooms r ON r.id=b.room_id WHERE b.group_id=? ORDER BY r.building, r.sort_order, r.name""",
            (booking["group_id"],)).fetchall()
    rooms = conn.execute("SELECT * FROM rooms WHERE active=1 AND out_of_order=0 ORDER BY CASE building WHEN 'Corp A' THEN 1 WHEN 'Corp B' THEN 2 WHEN 'Corp C' THEN 3 ELSE 4 END, sort_order, name").fetchall()
    if request.method == "POST":
        try:
            data = get_form_data(request.form)
            data["group_id"] = booking["group_id"]
            data["id"] = booking_id
        except (TypeError, ValueError):
            data = get_form_data_safe(request.form)
            data["group_id"] = booking["group_id"]
            data["id"] = booking_id
            flash("Verificați valorile numerice introduse.", "error")
            conn.close()
            return render_template("booking_form.html", rooms=rooms, data=data, mode="edit", booking_id=booking_id, calendar_start=data.get("check_in") or clock.today().isoformat(), group_bookings=group_bookings)
        valid_room_ids = {r["id"] for r in rooms}
        error = validate(data, valid_room_ids)
        if not error:
            selected_room = conn.execute("SELECT active, out_of_order FROM rooms WHERE id=?", (data["room_id"],)).fetchone()
            if not selected_room or not selected_room["active"] or selected_room["out_of_order"]:
                error = "Camera selectată este indisponibilă sau marcată ca defectă. Alegeți o cameră disponibilă."
        if not error and booking["group_id"] and data["status"] != "provisional" and not data.get("receipt_number", "").strip():
            error = "Rezervarea de grup nu poate fi confirmată fără numărul chitanței."
        if error:
            flash(error, "error")
            conn.close()
            return render_template("booking_form.html", rooms=rooms, data=data, mode="edit", booking_id=booking_id, calendar_start=data.get("check_in") or clock.today().isoformat(), group_bookings=group_bookings)
        conn.execute("BEGIN IMMEDIATE")
        conflict = overlaps(conn, data["room_id"], data["check_in"], data["check_out"], exclude_id=booking_id)
        if conflict:
            conn.rollback()
            flash(f"Conflict: camera este deja rezervată pentru {conflict['guest_name']} "
                  f"({conflict['check_in']} – {conflict['check_out']}).", "error")
            conn.close()
            return render_template("booking_form.html", rooms=rooms, data=data, mode="edit", booking_id=booking_id, calendar_start=data.get("check_in") or clock.today().isoformat(), group_bookings=group_bookings)
        conn.execute("""
            UPDATE bookings SET guest_name=?, phone=?, email=?, room_id=?, check_in=?, check_out=?,
            guests=?, rate=?, deposit=?, source=?, status=?, notes=?, late_arrival=?, night_note=?, receipt_number=?, updated_at=CURRENT_TIMESTAMP WHERE id=?
        """, tuple(data[k] for k in ["guest_name","phone","email","room_id","check_in","check_out","guests","rate","deposit","source","status","notes","late_arrival","night_note"]) + (data.get("receipt_number", ""), booking_id))
        apply_to_group = request.form.get("apply_to_group") == "1"
        if booking["group_id"] and apply_to_group:
            conn.execute("UPDATE bookings SET status=?, receipt_number=?, updated_at=CURRENT_TIMESTAMP WHERE group_id=? AND id != ?", (data["status"], data.get("receipt_number", ""), booking["group_id"], booking_id))
        conn.commit()
        conn.close()
        audit("update", "booking", booking_id, {"room_id": data["room_id"], "check_in": data["check_in"], "check_out": data["check_out"], "status": data["status"]})
        flash("Rezervarea și întregul grup au fost confirmate cu numărul chitanței." if booking["group_id"] and apply_to_group else "Rezervarea a fost actualizată.", "success")
        return redirect(url_for("index", start=data["check_in"]))
    data = dict(booking)
    conn.close()
    return render_template("booking_form.html", rooms=rooms, data=data, mode="edit", booking_id=booking_id, calendar_start=data.get("check_in") or clock.today().isoformat(), group_bookings=group_bookings)

@app.route("/booking/<int:booking_id>/reconfirm", methods=["POST"])
@login_required
def reconfirm_booking(booking_id):
    module_required("mod_reconfirm")
    conn = connect()
    booking = conn.execute("SELECT * FROM bookings WHERE id=?", (booking_id,)).fetchone()
    if not booking:
        conn.close(); abort(404)
    if booking['source'] != 'Direct' or booking['check_in'] != clock.today().isoformat() or booking['group_id']:
        conn.close(); flash("Reconfirmarea rapidă este disponibilă doar pentru rezervări directe individuale cu sosire astăzi.", "error")
        return redirect(url_for('edit_booking', booking_id=booking_id))
    conflict = overlaps(conn, booking['room_id'], booking['check_in'], booking['check_out'], exclude_id=booking_id)
    if conflict:
        conn.close(); flash("Camera a fost realocată între timp. Rezervarea nu poate fi reconfirmată fără rezolvarea conflictului.", "error")
        return redirect(url_for('edit_booking', booking_id=booking_id))
    conn.execute("UPDATE bookings SET status='confirmed', reconfirmed_at=?, updated_at=CURRENT_TIMESTAMP WHERE id=?", (clock.now().astimezone().isoformat(timespec='seconds'), booking_id))
    conn.close()
    audit("reconfirm", "booking", booking_id, {"source":"Direct", "check_in":booking['check_in']})
    flash("Rezervarea a fost reconfirmată.", "success")
    return redirect(url_for('edit_booking', booking_id=booking_id))


@app.route("/booking/group/<group_id>/cancel", methods=["POST"])
@login_required
def cancel_booking_group(group_id):
    """Anulare de grup: întregul grup SAU doar camerele bifate, numai după confirmare explicită în fereastra de dialog."""
    scope = request.form.get("scope", "")
    if scope not in {"all", "selected"} or request.form.get("confirmed") != "1":
        flash("Anularea nu a fost confirmată. Nu s-a modificat nimic.", "error")
        return redirect(url_for("index"))
    selected_ids = set()
    for value in request.form.getlist("booking_ids"):
        try:
            selected_ids.add(int(value))
        except ValueError:
            flash("Selecție invalidă. Nu s-a modificat nimic.", "error")
            return redirect(url_for("index"))
    conn = connect()
    rows = conn.execute("""SELECT b.id, r.name AS room_name FROM bookings b JOIN rooms r ON r.id=b.room_id
                           WHERE b.group_id=? AND b.status!='cancelled' ORDER BY r.sort_order, r.name""", (group_id,)).fetchall()
    if not rows:
        conn.close()
        flash("Nu există rezervări active pentru acest grup.", "error")
        return redirect(url_for("index"))
    active_ids = {r["id"] for r in rows}
    if scope == "selected":
        if not selected_ids:
            conn.close()
            flash("Bifați cel puțin o cameră pentru anulare. Nu s-a modificat nimic.", "error")
            return redirect(url_for("index"))
        if not selected_ids <= active_ids:       # id-uri străine sau deja anulate: refuzăm tot, nu anulăm parțial
            conn.close()
            flash("Selecția conține camere care nu aparțin acestui grup sau sunt deja anulate. Nu s-a modificat nimic.", "error")
            return redirect(url_for("index"))
        affected = [r for r in rows if r["id"] in selected_ids]
    else:
        affected = list(rows)
    ids = [r["id"] for r in affected]
    placeholders = ",".join("?" for _ in ids)
    conn.execute(f"UPDATE bookings SET status='cancelled', updated_at=CURRENT_TIMESTAMP WHERE id IN ({placeholders})", ids)
    conn.commit()
    conn.close()
    names = ", ".join(r["room_name"] for r in affected)
    audit("cancel_group", "booking_group", group_id, {"scope": scope, "booking_ids": ids, "rooms": names})
    remaining = len(rows) - len(ids)
    if remaining == 0:
        flash(f"Întregul grup a fost anulat ({len(ids)} camere: {names}).", "success")
    else:
        flash(f"Au fost anulate {len(ids)} camere ({names}). Rămân active {remaining} camere din grup.", "success")
    return redirect(url_for("index"))

@app.route("/booking/<int:booking_id>/cancel", methods=["POST"])
@login_required
def cancel_booking(booking_id):
    conn = connect()
    booking = conn.execute("SELECT id FROM bookings WHERE id=?", (booking_id,)).fetchone()
    if not booking:
        conn.close(); abort(404)
    conn.execute("UPDATE bookings SET status='cancelled', updated_at=CURRENT_TIMESTAMP WHERE id=?", (booking_id,))
    conn.close()
    audit("cancel", "booking", booking_id, {})
    flash("Rezervarea a fost anulată, iar camera este din nou disponibilă.", "success")
    return redirect(url_for("index"))

@app.route("/api/availability")
@login_required
def availability():
    start = request.args.get("start", clock.today().isoformat())
    end = request.args.get("end", (clock.today()+timedelta(days=14)).isoformat())
    conn = connect()
    conflicts = conn.execute("""
        SELECT r.name AS room, a.id AS booking_a, b.id AS booking_b
        FROM bookings a JOIN bookings b ON a.room_id=b.room_id AND a.id < b.id
        JOIN rooms r ON r.id=a.room_id
        WHERE a.status NOT IN ('cancelled', 'unconfirmed') AND b.status NOT IN ('cancelled', 'unconfirmed')
          AND a.check_in < b.check_out AND a.check_out > b.check_in
          AND a.check_in < ? AND a.check_out > ?
    """, (end, start)).fetchall()
    conn.close()
    return jsonify({"ok": len(conflicts) == 0, "conflicts": [dict(x) for x in conflicts]})

settings.connect = connect
init_db()
settings.invalidate()
try:
    os.chmod(DB_PATH, 0o600)      # baza de date poate fi citită doar de contul care rulează aplicația (pe Windows se ignoră)
except OSError:
    pass
integration_api.register(app, connect)

if __name__ == "__main__":
    # Local test scheduler: mark unreconfirmed same-day direct arrivals as unconfirmed at/after 20:30.
    scheduler = threading.Thread(target=cutoff_worker, daemon=True, name="direct-booking-cutoff")
    scheduler.start()
    # Local demo only. For internet deployment, use a production WSGI server and HTTPS.
    host = os.environ.get("MOTEL_HOST", "127.0.0.1")
    if host not in ("127.0.0.1", "localhost") and len(app.secret_key) < 32:
        raise SystemExit("Pentru acces din rețea, MOTEL_SECRET_KEY trebuie să aibă minimum 32 de caractere (sau lăsați aplicația să genereze automat cheia).")
    cert, key = os.environ.get("MOTEL_SSL_CERT"), os.environ.get("MOTEL_SSL_KEY")
    port = int(os.environ.get("MOTEL_PORT", "5000"))
    scheme = "https" if cert and key else "http"
    if clock.is_test():
        print(f"ATENȚIE: CEAS DE TEST activ ({clock.now():%d.%m.%Y %H:%M}). Nu folosiți date reale.")
    if os.environ.get("MOTEL_SERVER", "auto") != "dev":
        try:
            from cheroot import wsgi as cheroot_wsgi
        except ImportError:
            cheroot_wsgi = None
            print("AVERTISMENT: pachetul „cheroot” lipsește; pornesc serverul de dezvoltare. Rulați: pip install -r requirements.txt")
        if cheroot_wsgi is not None:
            server = cheroot_wsgi.Server((host, port), app, numthreads=16, server_name="PMS")
            server.version = "PMS"                                   # antetul Server nu dezvăluie versiuni
            if cert and key:
                from cheroot.ssl.builtin import BuiltinSSLAdapter
                server.ssl_adapter = BuiltinSSLAdapter(cert, key)
            print(f"PMS pornit pe serverul de producție: {scheme}://{host}:{port}  (oprire: Ctrl+C)")
            try:
                server.start()
            except KeyboardInterrupt:
                server.stop()
            raise SystemExit(0)
    from werkzeug.serving import WSGIRequestHandler
    WSGIRequestHandler.version_string = lambda self: "PMS"   # nu dezvăluim versiunile Werkzeug/Python în antetul Server
    ssl_context = (cert, key) if cert and key else None
    print(f"PMS pornit pe serverul de dezvoltare: {scheme}://{host}:{port}")
    app.run(host=host, port=port, debug=False, ssl_context=ssl_context, threaded=True)
